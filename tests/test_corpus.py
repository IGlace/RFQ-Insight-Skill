"""The corpus: folder walk, per-file extraction, the checkpoint on disk, and
the delimited text the model actually reads.

Ported from the old test_extract_corpus.py (extract_corpus CLI) and
test_corpus_text.py (corpus_text CLI); both now live in rfq/corpus.py.
"""

import io
import json
import os
import re

from rfq import corpus as cp
from rfq.extractors import marker_for
from rfq.figures import filter_images, write_images

SKILL = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                     "plugins", "rfq-insight", "skills", "rfq-insight")
PROMPTS = os.path.join(SKILL, "references", "prompts")


def _docx(path, paragraph="The system shall deliver 450 Nm."):
    import docx
    d = docx.Document()
    d.add_paragraph(paragraph)
    d.save(path)


def _noise_png(w, h):
    """A PNG big enough in both bytes and pixels to survive the filter."""
    from PIL import Image
    buf = io.BytesIO()
    Image.frombytes("RGB", (w, h), os.urandom(w * h * 3)).save(buf, format="PNG")
    return buf.getvalue()


def _docx_with_picture(path, png_bytes, paragraph="Body text."):
    import docx
    from docx.shared import Inches
    png = os.path.join(os.path.dirname(path), "_tmp_pic.png")
    with open(png, "wb") as fh:
        fh.write(png_bytes)
    d = docx.Document()
    d.add_paragraph(paragraph)
    d.add_picture(png, width=Inches(2))
    d.save(path)
    os.remove(png)


def _rfq_folder(tmp_path):
    folder = tmp_path / "Acme V6"
    (folder / "annexes").mkdir(parents=True)
    _docx(str(folder / "SOR.docx"))
    (folder / "notes.txt").write_text("Quotation due 20.06.2026",
                                      encoding="utf-8")
    _docx(str(folder / "annexes" / "Annex A.docx"), "Annex content.")
    (folder / "broken.docx").write_bytes(b"not a docx")
    (folder / "onenote.one").write_bytes(b"\x00\x01")
    return str(folder)


# ============================================================== folder walk
def test_walk_finds_supported_files_recursively_and_sorted(tmp_path):
    folder = _rfq_folder(tmp_path)
    supported, ignored = cp.walk_supported(folder)
    names = [os.path.basename(p) for p in supported]
    assert names == ["Annex A.docx", "SOR.docx", "broken.docx", "notes.txt"]
    assert ignored == ["onenote.one"]


def test_walk_never_descends_into_our_own_output_folder(tmp_path):
    folder = _rfq_folder(tmp_path)
    prior = os.path.join(folder, "output")
    os.makedirs(prior)
    _docx(os.path.join(prior, "Previous_RFQ_Analysis.docx"))
    supported, _ = cp.walk_supported(folder)
    assert all("output" not in os.path.relpath(p, folder).split(os.sep)
               for p in supported)


def test_walk_skips_dot_directories(tmp_path):
    folder = _rfq_folder(tmp_path)
    hidden = os.path.join(folder, ".git")
    os.makedirs(hidden)
    _docx(os.path.join(hidden, "Hidden.docx"))
    supported, ignored = cp.walk_supported(folder)
    assert not any("Hidden.docx" in p for p in supported)
    assert "Hidden.docx" not in ignored


def test_the_work_dir_sits_under_output(tmp_path):
    assert cp.work_dir(str(tmp_path)) == os.path.join(
        str(tmp_path), "output", ".work")


# ============================================================= build_corpus
def test_build_corpus_carries_the_per_file_contract(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.build_corpus(folder, "Acme V6")

    assert corpus["projectName"] == "Acme V6"
    assert corpus["oppId"] == ""
    by_name = {f["filename"]: f for f in corpus["filesContents"]}
    assert set(by_name) == {"SOR.docx", "notes.txt", "Annex A.docx"}

    sor = by_name["SOR.docx"]
    assert sor["extension"] == "docx"
    assert sor["category"] == "office_doc"
    assert sor["extraction_method"] == "python-docx"
    assert sor["path"] == "SOR.docx"
    assert sor["file_size"] > 0
    assert "450 Nm" in sor["extracted_text"]
    assert by_name["Annex A.docx"]["path"] == os.path.join("annexes",
                                                           "Annex A.docx")


def test_every_file_carries_the_sha256_of_its_bytes(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.build_corpus(folder, "Acme V6")
    for entry in corpus["filesContents"]:
        path = os.path.join(folder, entry["path"])
        assert entry["sha256"] == cp.file_sha256(path)
        assert re.fullmatch(r"[0-9a-f]{64}", entry["sha256"])


def test_unreadable_file_is_recorded_with_a_reason_not_dropped(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.build_corpus(folder, "Acme V6")
    skipped = corpus["extraction"]["skipped"]
    assert [s["filename"] for s in skipped] == ["broken.docx"]
    assert skipped[0]["path"] == "broken.docx"
    assert "corrupt" in skipped[0]["reason"]
    assert corpus["extraction"]["ignored"] == ["onenote.one"]
    assert corpus["extraction"]["parsed"] == 3


def test_preprocess_stats_are_recorded_per_file(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.build_corpus(folder, "Acme V6")
    stats = corpus["extraction"]["preprocess"]
    assert set(stats) == {"SOR.docx", "notes.txt", "Annex A.docx"}
    assert stats["SOR.docx"]["chars_after"] > 0


def test_an_empty_folder_builds_an_empty_corpus(tmp_path):
    folder = tmp_path / "Empty"
    folder.mkdir()
    corpus, raw = cp.build_corpus(str(folder), "Empty")
    assert corpus["filesContents"] == [] and raw == []
    assert corpus["extraction"]["parsed"] == 0


# =========================================================== run_extraction
def test_run_extraction_writes_the_three_checkpoints(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, manifest = cp.run_extraction(folder, "Acme V6")

    work = cp.work_dir(folder)
    with open(os.path.join(work, "corpus.json"), encoding="utf-8") as fh:
        assert json.load(fh) == corpus
    with open(os.path.join(work, "images.json"), encoding="utf-8") as fh:
        assert json.load(fh) == manifest

    assert corpus["projectName"] == "Acme V6"
    assert os.path.isdir(os.path.join(work, "images"))
    assert manifest["images"] == []          # nothing embedded in the fixtures
    assert corpus["extraction"]["images_flagged"] == []


def test_project_name_defaults_to_the_folder_name(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.run_extraction(folder)
    assert corpus["projectName"] == "Acme V6"
    with open(os.path.join(cp.work_dir(folder), "corpus.json"),
              encoding="utf-8") as fh:
        assert json.load(fh)["projectName"] == "Acme V6"


def test_summary_is_plain_language(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, raw = cp.build_corpus(folder, "Acme V6")
    kept, dropped = filter_images(raw)
    manifest = write_images(kept, dropped,
                            os.path.join(cp.work_dir(folder), "images"))
    text = "\n".join(cp.summary_lines(corpus, manifest))

    assert "Read 3 of 4 files" in text
    assert "broken.docx" in text and "corrupt" in text
    assert "onenote.one" in text
    for jargon in ("corpus", "extraction_method", "Traceback", "None"):
        assert jargon not in text


# ------------------------------------------- image identity across the corpus
def _colliding_basenames(tmp_path):
    """annexes/Spec.docx and attachments/Spec.docx, each with its OWN
    picture."""
    folder = tmp_path / "Acme V6"
    (folder / "annexes").mkdir(parents=True)
    (folder / "attachments").mkdir(parents=True)
    annex_png, attach_png = _noise_png(300, 300), _noise_png(301, 301)
    _docx_with_picture(str(folder / "annexes" / "Spec.docx"), annex_png,
                       "Annex body.")
    _docx_with_picture(str(folder / "attachments" / "Spec.docx"), attach_png,
                       "Attachment body.")
    return str(folder), annex_png, attach_png


def test_same_named_documents_get_distinct_image_ids(tmp_path):
    folder, _, _ = _colliding_basenames(tmp_path)
    _, raw = cp.build_corpus(folder, "Acme V6")

    ids = [i.image_id for i in raw]
    assert len(ids) == 2
    assert len(set(ids)) == 2, f"colliding image ids: {ids}"
    assert sorted(ids) == ["annexes_Spec.docx__img01",
                           "attachments_Spec.docx__img01"]
    # `source` stays the display name - it is what the user is shown.
    assert {i.source for i in raw} == {"Spec.docx"}


def test_same_named_documents_keep_their_own_picture_on_disk(tmp_path):
    folder, annex_png, attach_png = _colliding_basenames(tmp_path)
    _, manifest = cp.run_extraction(folder, "Acme V6")

    work = cp.work_dir(folder)
    ids = [e["id"] for e in manifest["images"]]
    assert sorted(ids) == ["annexes_Spec.docx__img01",
                           "attachments_Spec.docx__img01"]
    files = sorted(os.listdir(os.path.join(work, "images")))
    assert files == ["annexes_Spec.docx__img01.png",
                     "attachments_Spec.docx__img01.png"]

    on_disk = {}
    for name in files:
        with open(os.path.join(work, "images", name), "rb") as fh:
            on_disk[name] = fh.read()
    assert on_disk["annexes_Spec.docx__img01.png"] == annex_png
    assert on_disk["attachments_Spec.docx__img01.png"] == attach_png


def test_each_same_named_document_carries_only_its_own_marker(tmp_path):
    folder, _, _ = _colliding_basenames(tmp_path)
    corpus, _ = cp.run_extraction(folder, "Acme V6")
    by_path = {f["path"]: f["extracted_text"] for f in corpus["filesContents"]}

    annex = by_path[os.path.join("annexes", "Spec.docx")]
    attach = by_path[os.path.join("attachments", "Spec.docx")]
    assert marker_for("annexes_Spec.docx__img01") in annex
    assert marker_for("annexes_Spec.docx__img01") not in attach
    assert marker_for("attachments_Spec.docx__img01") in attach
    assert marker_for("attachments_Spec.docx__img01") not in annex


def test_identical_copies_of_a_document_still_get_their_own_id(tmp_path):
    folder = tmp_path / "Acme V6"
    (folder / "annexes").mkdir(parents=True)
    (folder / "attachments").mkdir(parents=True)
    png = _noise_png(300, 300)
    _docx_with_picture(str(folder / "annexes" / "Spec.docx"), png, "Body.")
    _docx_with_picture(str(folder / "attachments" / "Spec.docx"), png, "Body.")

    _, raw = cp.build_corpus(str(folder), "Acme V6")
    assert sorted(i.image_id for i in raw) == ["annexes_Spec.docx__img01",
                                               "attachments_Spec.docx__img01"]


def test_a_top_level_file_keeps_its_plain_name_in_the_image_id(tmp_path):
    folder = tmp_path / "Acme V6"
    folder.mkdir()
    _docx_with_picture(str(folder / "SOR.docx"), _noise_png(300, 300))
    _, raw = cp.build_corpus(str(folder), "Acme V6")
    assert [i.image_id for i in raw] == ["SOR.docx__img01"]


def test_two_paths_that_slug_alike_are_still_told_apart(tmp_path):
    folder = tmp_path / "Acme V6"
    (folder / "a b").mkdir(parents=True)
    (folder / "a_b").mkdir(parents=True)
    _docx_with_picture(str(folder / "a b" / "S.docx"), _noise_png(300, 300))
    _docx_with_picture(str(folder / "a_b" / "S.docx"), _noise_png(301, 301))

    _, raw = cp.build_corpus(str(folder), "Acme V6")
    ids = [i.image_id for i in raw]
    assert len(set(ids)) == 2, f"colliding image ids: {ids}"


# ------------------------------------------ no marker without a picture behind
def test_a_filtered_out_picture_leaves_no_transcription_marker(tmp_path):
    folder = tmp_path / "Acme V6"
    folder.mkdir()
    _docx_with_picture(str(folder / "SOR.docx"), _noise_png(16, 16),
                       "Requirement text.")

    corpus, manifest = cp.run_extraction(str(folder), "Acme V6")
    assert manifest["images"] == []
    assert [d["reason"] for d in manifest["filtered"]] == ["too_small"]
    assert corpus["extraction"]["images_filtered"] == manifest["filtered"]

    text = corpus["filesContents"][0]["extracted_text"]
    assert "transcription pending" not in text
    assert marker_for("SOR.docx__img01") not in text
    assert "Requirement text." in text
    assert corpus["extraction"]["chars"] == len(text)


def test_a_kept_picture_keeps_its_marker(tmp_path):
    folder = tmp_path / "Acme V6"
    folder.mkdir()
    _docx_with_picture(str(folder / "SOR.docx"), _noise_png(300, 300))

    corpus, manifest = cp.run_extraction(str(folder), "Acme V6")
    text = corpus["filesContents"][0]["extracted_text"]
    assert marker_for("SOR.docx__img01") in text
    assert corpus["extraction"]["images_flagged"] == ["SOR.docx__img01"]
    assert manifest["images"][0]["sha256"]


def test_dropping_a_marker_leaves_the_surrounding_text_intact():
    marker = marker_for("SOR.docx__img01")
    corpus = {"filesContents": [
        {"filename": "SOR.docx",
         "extracted_text": f"Before.\n\n{marker}\n\nAfter."}],
        "extraction": {"chars": 0}}
    assert cp.drop_markers(corpus, [{"id": "SOR.docx__img01"}]) == 1
    assert corpus["filesContents"][0]["extracted_text"] == "Before.\n\nAfter."
    assert corpus["extraction"]["chars"] == len("Before.\n\nAfter.")


def test_the_skipped_picture_summary_names_the_real_reason():
    filtered = [
        {"id": "a", "reason": "too_small", "explanation": "x"},
        {"id": "b", "reason": "cap_reached", "explanation": "y"},
        {"id": "c", "reason": "unreadable", "explanation": "z"},
    ]
    lines = cp.filtered_lines(filtered)
    text = "\n".join(lines)
    assert len(lines) == 3
    assert "1 picture(s) that are too small" in text
    assert "beyond the limit" in text
    assert "could not open" in text
    for jargon in ("cap_reached", "too_small", "unreadable"):
        assert jargon not in text


# ============================================================ format_corpus
def _files():
    return [
        {"filename": "SOR.docx", "extension": "docx", "path": "SOR.docx",
         "extracted_text": "The system shall deliver 450 Nm."},
        {"filename": "Annex A.pdf", "extension": "pdf",
         "path": os.path.join("annexes", "Annex A.pdf"),
         "extracted_text": "Quotation due 20.06.2026"},
    ]


def _corpus_folder(tmp_path, files=None):
    work = tmp_path / "Acme V6" / "output" / ".work"
    work.mkdir(parents=True)
    with open(work / "corpus.json", "w", encoding="utf-8") as fh:
        json.dump({"projectName": "Acme V6",
                   "filesContents": files if files is not None else _files()},
                  fh)
    return str(tmp_path / "Acme V6")


def test_each_file_becomes_a_delimited_block():
    text = cp.format_corpus(_files())
    assert "=== FILE: SOR.docx ===" in text
    assert "=== FILE: Annex A.pdf ===" in text
    assert "(type: docx; path: SOR.docx)" in text
    assert "The system shall deliver 450 Nm." in text
    assert "Quotation due 20.06.2026" in text


def test_the_header_matches_the_contract_the_prompt_states():
    """The prompt tells the model the delimiter is a LINE of the form
    `=== FILE: <name> ===` and that source filenames are copied exactly
    from it."""
    text = cp.format_corpus(_files())
    headers = re.findall(r"(?m)^=== FILE: (.+) ===$", text)
    assert headers == ["SOR.docx", "Annex A.pdf"]
    assert headers == [f["filename"] for f in _files()]


def test_the_block_order_and_shape_are_exactly_the_container_app_s():
    text = cp.format_corpus(_files())
    assert text == (
        "=== FILE: SOR.docx ===\n"
        "(type: docx; path: SOR.docx)\n"
        "The system shall deliver 450 Nm.\n"
        "\n"
        "=== FILE: Annex A.pdf ===\n"
        f"(type: pdf; path: {os.path.join('annexes', 'Annex A.pdf')})\n"
        "Quotation due 20.06.2026")


def test_the_prompts_still_describe_the_format_this_module_emits():
    """If someone edits the delimiter here or there, these stop agreeing."""
    prompts = [n for n in os.listdir(PROMPTS)
               if n.startswith("RFQ_QA_Extraction_Prompt_")]
    assert prompts, "the QA extraction prompt is missing"
    for name in prompts:
        with open(os.path.join(PROMPTS, name), encoding="utf-8") as fh:
            assert "=== FILE: <name> ===" in fh.read()
    sample = cp.format_corpus([{"filename": "SOR.docx", "extension": "docx",
                                "path": "SOR.docx", "extracted_text": "x"}])
    assert sample.startswith("=== FILE: ")
    assert sample.splitlines()[0].endswith(" ===")


def test_an_empty_corpus_is_an_empty_string():
    assert cp.format_corpus([]) == ""


def test_a_merged_transcription_travels_inside_its_parent_block():
    files = [{"filename": "SOR.docx", "extension": "docx", "path": "SOR.docx",
              "extracted_text": "Intro.\n\n[BEGIN IMAGE TRANSCRIPTION "
                                "SOR.docx__img01]\nSOP Q3 2027.\n"
                                "[END IMAGE TRANSCRIPTION]"},
             {"filename": "Notes.txt", "extension": "txt", "path": "Notes.txt",
              "extracted_text": "Unrelated."}]
    text = cp.format_corpus(files)
    sor_block = text.split("=== FILE: Notes.txt ===")[0]
    assert "SOP Q3 2027." in sor_block


def test_long_lines_are_wrapped_below_the_file_reader_limit_without_losing_text():
    long_paragraph = " ".join(f"word{i}" for i in range(900))
    wide_row = "| " + " | ".join(["cell content " * 20] * 12) + " |"
    files = [{"filename": "SOR.docx", "extension": "docx", "path": "SOR.docx",
              "extracted_text": long_paragraph + "\n\n" + wide_row}]
    text = cp.format_corpus(files)
    assert max(len(line) for line in text.splitlines()) <= cp.MAX_LINE_CHARS
    # Only line breaks were added: joining on whitespace restores every token.
    assert " ".join(text.split()) == " ".join(
        ("=== FILE: SOR.docx === (type: docx; path: SOR.docx) "
         + long_paragraph + " " + wide_row).split())


def test_the_soft_wrap_limit_is_fifteen_hundred_characters():
    assert cp.MAX_LINE_CHARS == 1_500


def test_short_lines_are_untouched_by_wrapping():
    body = "short line\n| a | b |\n\nanother"
    assert cp.soft_wrap(body) == body


def test_a_single_overlong_token_is_still_cut_rather_than_kept_whole():
    assert max(len(l) for l in cp.soft_wrap("x" * 4000).splitlines()) <= cp.MAX_LINE_CHARS


# ======================================================== write_corpus_text
def test_write_corpus_text_writes_corpus_txt_in_the_work_dir(tmp_path):
    folder = _corpus_folder(tmp_path)
    out = cp.write_corpus_text(folder)
    assert out == os.path.join(folder, "output", ".work", "corpus.txt")
    with open(out, encoding="utf-8") as fh:
        assert fh.read() == cp.format_corpus(_files())


def test_write_corpus_text_honours_an_explicit_path(tmp_path):
    folder = _corpus_folder(tmp_path)
    out_path = str(tmp_path / "out" / "corpus.txt")
    assert cp.write_corpus_text(folder, out_path) == out_path
    with open(out_path, encoding="utf-8") as fh:
        assert fh.read() == cp.format_corpus(_files())


def test_write_corpus_text_also_writes_the_source_catalogue(tmp_path):
    folder = _corpus_folder(tmp_path)
    cp.write_corpus_text(folder)
    with open(os.path.join(folder, "output", ".work", "source_catalog.json"),
              encoding="utf-8") as fh:
        catalogue = json.load(fh)
    assert catalogue["version"] == 1
    assert {r["filename"] for r in catalogue["references"]} == {
        "SOR.docx", "Annex A.pdf"}
    for row in catalogue["references"]:
        assert row["id"].startswith("src-")


def test_write_corpus_text_raises_when_there_is_no_corpus(tmp_path):
    folder = tmp_path / "Acme V6"
    folder.mkdir()
    try:
        cp.write_corpus_text(str(folder))
    except FileNotFoundError:
        pass
    else:
        raise AssertionError("a missing corpus.json must be reported")


# ========================================================== corpus_revision
def _revision_corpus(text="Body.", sha="a" * 64, skipped=None):
    return {"projectName": "Demo", "filesContents": [
        {"filename": "SOR.docx", "path": "SOR.docx", "sha256": sha,
         "extracted_text": text, "modified": "2026-01-01T00:00:00",
         "processed_at": "2026-01-01T00:00:00"}],
        "extraction": {"skipped": skipped or []}}


def test_corpus_revision_is_stable_and_ignores_timestamps():
    a = _revision_corpus()
    b = _revision_corpus()
    b["filesContents"][0]["processed_at"] = "2030-01-01T00:00:00"
    b["filesContents"][0]["modified"] = "2030-01-01T00:00:00"
    b["projectName"] = "Renamed"
    assert cp.corpus_revision(a) == cp.corpus_revision(b)
    assert re.fullmatch(r"[0-9a-f]{64}", cp.corpus_revision(a))


def test_corpus_revision_tracks_text_bytes_and_skipped_inputs():
    base = cp.corpus_revision(_revision_corpus())
    assert cp.corpus_revision(_revision_corpus(text="Other.")) != base
    assert cp.corpus_revision(_revision_corpus(sha="b" * 64)) != base
    assert cp.corpus_revision(_revision_corpus(
        skipped=[{"filename": "x.pdf", "path": "x.pdf", "reason": "r",
                  "detail": ""}])) != base


def test_corpus_revision_does_not_depend_on_file_order():
    a = {"filesContents": [{"filename": "A", "path": "A", "sha256": "1",
                            "extracted_text": "a"},
                           {"filename": "B", "path": "B", "sha256": "2",
                            "extracted_text": "b"}], "extraction": {}}
    b = {"filesContents": list(reversed(a["filesContents"])), "extraction": {}}
    assert cp.corpus_revision(a) == cp.corpus_revision(b)


# =========================================================== corpus_changes
def test_a_fresh_extraction_reports_no_changes(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.build_corpus(folder, "Acme V6")
    assert cp.corpus_changes(folder, corpus) == {
        "added": [], "removed": [], "changed": []}


def test_corpus_changes_names_added_removed_and_changed_files(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.build_corpus(folder, "Acme V6")

    (tmp_path / "Acme V6" / "new.txt").write_text("New.", encoding="utf-8")
    os.remove(os.path.join(folder, "notes.txt"))
    _docx(os.path.join(folder, "SOR.docx"), "Rewritten.")

    assert cp.corpus_changes(folder, corpus) == {
        "added": ["new.txt"], "removed": ["notes.txt"], "changed": ["SOR.docx"]}


def test_a_skipped_file_is_not_reported_as_added(tmp_path):
    """broken.docx was seen and recorded as unreadable; it is not new."""
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.build_corpus(folder, "Acme V6")
    assert "broken.docx" not in cp.corpus_changes(folder, corpus)["added"]


def test_a_file_without_a_recorded_hash_counts_as_changed(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = cp.build_corpus(folder, "Acme V6")
    for entry in corpus["filesContents"]:
        if entry["path"] == "SOR.docx":
            del entry["sha256"]
    assert cp.corpus_changes(folder, corpus)["changed"] == ["SOR.docx"]
