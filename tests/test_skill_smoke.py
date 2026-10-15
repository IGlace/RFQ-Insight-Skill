"""Run the installed plugin's runner end to end with canned worker outputs.

The plugin folder is copied to an unrelated location first, the way a host
installs it, and `rfq_run.py next` is driven by subprocess from a different
working directory. The test plays the coordinator: it dispatches every
request the runner returns by writing the worker's output files itself and
records the worker lifecycle with `event`. This checks relocation, the
filesystem contracts, checkpoints, resume and restart; it does not certify
live model judgement.
"""

import json
import os
from pathlib import Path
import shutil
import subprocess
import sys

from PIL import Image, ImageDraw
import pytest

from conftest import REPO_ROOT
from rfq.schema import TITLES

ROOT = Path(REPO_ROOT)
PLUGIN = ROOT / "plugins" / "rfq-insight"
VOCAB = json.loads((PLUGIN / "skills" / "rfq-insight" / "references" / "data"
                    / "vocabularies.json").read_text(encoding="utf-8"))

SOW = "01_statement_of_work.txt"
APPENDIX = "02_appendix_b.txt"
TRANSCRIPTION = "Review: 19 December 2026"


# ------------------------------------------------------- canned worker output
def _opportunity() -> dict:
    """Section 1 with official Guidance vocabulary values."""
    product = VOCAB["products"][3]["name"]
    return {
        "customer": "Northwind Mobility", "customer_site": "Info not available",
        "customer_contact": "Dana Example",
        "quotation_due_date": "20 October 2026",
        "customer_short": "Northwind", "project_title": "Cedar Controller",
        "opportunity_name": "Northwind – Cedar Controller",
        "account_name": "Info not available", "kam_pte": "Info not available",
        "project_type": VOCAB["project_types"][3]["name"],
        "lead_product": product,
        "technology": VOCAB["technologies"][3]["name"],
        "industry": VOCAB["industries"][0]["name"],
        "products": [{"name": product, "percent": 100}],
        "skills": [{"name": VOCAB["skills"][3]["name"], "percent": 100}],
    }


# Every answer has its section's canonical markdown shape (CONTRACT §9),
# so the runner plans the slides itself and no slide worker is needed.
_ANSWERS = {
    1: ("| Field | Value |\n|---|---|\n| Customer | Northwind Mobility |\n"
        "| Customer contact | Dana Example |\n"
        "| Quotation due date | 20 October 2026 |"),
    2: ("Northwind Mobility requests a proposal from AVL for the Cedar "
        "controller programme. The customer will provide prototype hardware."),
    3: ("| Item | Value |\n|---|---|\n| Sensor modules | four |\n"
        "| Supply voltage | 9 to 16 V |\n| Production volume | not stated |"),
    4: ("| Application | Approach |\n|---|---|\n"
        "| Bench verification | ETS-42:2024 on customer prototype hardware |"),
    5: ("Hub: Northwind\n\n| Party | Type | Role |\n|---|---|---|\n"
        "| Northwind Mobility | customer | Provides the prototype hardware |\n"
        "| AVL | avl | Delivers the engineering work |\n\n"
        "No further parties are named in the RFQ."),
    6: ("| Phase | Definition | Scope | Hardware milestone | Span |\n"
        "|---|---|---|---|---|\n"
        "| Bench Readiness | No definition in the RFQ | in scope | | "
        "2 November 2026 – 27 November 2026 |\n"
        "| Vehicle Check | No definition in the RFQ | in scope | | "
        "30 November 2026 – 18 December 2026 |"),
    7: ("| Scope of Work | Bench Readiness | Vehicle Check |\n|---|---|---|\n"
        "| Requirements review | AVL | AVL |\n"
        "| Vehicle integration | Northwind | Northwind |\n"
        "| Interface approval | Joint | Joint |\n"
        "| Prototype shipping | — | — |"),
    8: ("| Deliverable | Development Phase / Task |\n|---|---|\n"
        "| Requirements review | Bench Readiness |\n"
        "| Bench test report | Vehicle Check |"),
    9: ("| Milestone | Target timing | Phase / Task |\n|---|---|---|\n"
        "| Bench Readiness | 2 November 2026 | Bench Readiness |\n"
        "| Vehicle Check | 30 November 2026 | Vehicle Check |\n"
        "| Review | 19 December 2026 | Vehicle Check |"),
    10: "| Standard | Applies to |\n|---|---|\n| ETS-42:2024 | Bench verification |",
    11: ("## Project Risks\n- Discrepancies between the statement and Appendix B "
         "must be clarified before quotation.\n\n"
         "## Document Risks\n- Production volume is not stated."),
}
# A section-4 table without its separator row: the runner cannot parse it,
# so the slide worker plans the whole deck.
UNPARSABLE_4 = "| Application | Approach |\n| Bench verification | ETS-42:2024 |"


def _analysis(answers=None) -> dict:
    """The model's sections 1-11, with sources naming real corpus files
    (Source Material is the runner's)."""
    answers = {**_ANSWERS, **(answers or {})}
    return {"answers": [
        {"id": i, "answer": answers[i], "not_found": False, "deductions": [],
         "sources": ([{"filename": SOW, "locator": ""}]
                     + ([{"filename": APPENDIX, "locator": ""}]
                        if i == 11 else []))}
        for i in range(1, 12)], "opportunity": _opportunity()}


def _table(number, columns, rows) -> dict:
    return {"layout": "table", "title": TITLES[number][0],
            "columns": columns, "rows": rows}


def _slide_plan(project: str) -> dict:
    """A draft plan with exactly the shapes schema.SECTION_SHAPES allows."""
    slides = {
        1: [{"layout": "opportunity", "title": TITLES[1][0]}],
        2: [{"layout": "prose", "title": TITLES[2][0],
             "body": "Northwind Mobility requests a proposal from AVL for "
                     "the Cedar controller programme."}],
        3: [_table(3, ["Item", "Value"], [["Sensor modules", "four"],
                                          ["Supply voltage", "9 to 16 V"]])],
        4: [_table(4, ["Approach", "Detail"],
                   [["Bench verification", "ETS-42:2024 on customer hardware"]])],
        5: [{"layout": "collab", "title": TITLES[5][0], "parties": [
            {"name": "Northwind Mobility", "role": "Customer", "type": "customer"},
            {"name": "AVL", "role": "Engineering partner", "type": "avl"}]}],
        6: [{"layout": "phases", "title": TITLES[6][0], "phases": [
            {"name": "Bench Readiness", "desc": "2 Nov - 27 Nov 2026"},
            {"name": "Vehicle Check", "desc": "30 Nov - 18 Dec 2026"}]}],
        7: [{"layout": "ownership", "title": TITLES[7][0],
             "customerAbbr": "Northwind",
             "columns": ["Scope of Work", "Bench Readiness", "Vehicle Check"],
             "rows": [["Requirements review", "AVL", "AVL"],
                      ["Vehicle integration", "Customer", "Customer"],
                      ["Interface approval", "Joint", "Joint"]]}],
        8: [_table(8, ["Deliverable", "Phase"],
                   [["Requirements review", "Bench Readiness"],
                    ["Bench test report", "Vehicle Check"]])],
        9: [{"layout": "gantt", "title": TITLES[9][0], "steps": [
                {"label": "Bench Readiness start", "date": "2026-11-02", "desc": ""},
                {"label": "Vehicle Check start", "date": "2026-11-30", "desc": ""},
                {"label": "Review", "date": "2026-12-19", "desc": ""}],
             "bars": [{"name": "Bench Readiness", "start": "2026-11-02",
                       "end": "2026-11-27"},
                      {"name": "Vehicle Check", "start": "2026-11-30",
                       "end": "2026-12-18"}]},
            dict(_table(9, ["Milestone", "Target timing", "Phase / Task"],
                        [["Bench Readiness start", "2026-11-02", "Bench Readiness"],
                         ["Vehicle Check start", "2026-11-30", "Vehicle Check"],
                         ["Review", "2026-12-19", "Review"]]),
                 title=TITLES[9][0] + " — Milestone Table")],
        10: [_table(10, ["Standard", "Scope"],
                    [["ETS-42:2024", "Bench verification"]])],
        11: [{"layout": "panel", "title": TITLES[11][0], "heading": "Project Risks",
              "items": ["Discrepancies between the statement and Appendix B "
                        "must be clarified."]},
             {"layout": "panel", "title": TITLES[11][0], "heading": "Document Risks",
              "items": ["Production volume is not stated."]}],
        12: [_table(12, ["Document", "Content"],
                    [[SOW, "Statement of work"], [APPENDIX, "Customer requirements"]])],
    }
    return {"projectName": project, "questions": [
        {"number": n, "titleEn": TITLES[n][0], "titleDe": TITLES[n][1],
         "slides": slides[n]} for n in range(1, 13)]}


def _serve(request: dict, answers=None) -> None:
    """Write the canned output(s) of one worker request; `answers`
    replaces some of the analysis answers."""
    stage = request["stage"]
    if stage == "figures":
        for out in request["outputs"]:
            path = Path(out)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(TRANSCRIPTION, encoding="utf-8")
        return
    out = Path(request["outputs"][0])
    out.parent.mkdir(parents=True, exist_ok=True)
    if stage == "analysis":
        payload = _analysis(answers)
    elif stage == "verification":
        payload = {"verdicts": [
            {"id": i, "verdict": "confirmed",
             "evidence": "Production volume is not stated."}
            for i in request["claim_ids"]]}
    elif stage == "slides":
        payload = _slide_plan("Cedar Demo")
    else:
        raise AssertionError(f"unexpected stage {stage}")
    out.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


# ------------------------------------------------------------- the driver
class Coordinator:
    """Drives the installed runner from an unrelated cwd, as a host would."""

    def __init__(self, scripts: Path, folder: Path, cwd: Path):
        self.scripts, self.folder, self.cwd = scripts, folder, cwd
        self.served = []          # (stage, request_id) in dispatch order
        self.answers = None       # analysis answers to serve instead
        self.completed = []       # request ids for the next `next` call

    def run(self, *args) -> dict:
        env = os.environ.copy()
        # Retain only installed dependencies, never source-repo imports.
        env["PYTHONPATH"] = os.pathsep.join(
            p for p in sys.path if Path(p).name == "site-packages")
        env["PYTHONIOENCODING"] = "utf-8"
        env["PYTHONDONTWRITEBYTECODE"] = "1"
        result = subprocess.run(
            [sys.executable, str(self.scripts / "rfq_run.py"), *map(str, args)],
            cwd=self.cwd, env=env, capture_output=True, text=True,
            encoding="utf-8", timeout=120)
        assert result.returncode == 0, result.stdout + result.stderr
        assert result.stderr == "", result.stderr
        return json.loads(result.stdout)

    def next(self, *extra) -> dict:
        """`next`, reporting every worker served since the last call with
        `--completed` (SKILL.md: the terminal observation rides along)."""
        done = [arg for request_id in self.completed
                for arg in ("--completed", request_id)]
        self.completed = []
        return self.run("next", self.folder, "--host", "claude",
                        "--project", "Cedar Demo", *done, *extra)

    def dispatch(self, result: dict) -> None:
        assert result["status"] in ("dispatch", "repair"), result
        for request in result["requests"]:
            assert request["host"] == "claude"
            assert request["inherit_history"] is False
            assert request["agent"] == "general-purpose"
            event = self.run("event", self.folder, "--request-id",
                             request["request_id"], "--event", "started")
            assert event["status"] == "recorded"
            _serve(request, self.answers)
            self.completed.append(request["request_id"])
            self.served.append((request["stage"], request["request_id"]))

    def loop(self, *extra, limit: int = 12) -> dict:
        """`next` until the runner stops needing workers."""
        for _ in range(limit):
            result = self.next(*extra)
            if result["status"] in ("done", "stopped", "blocked"):
                return result
            assert result["status"] == "dispatch", result
            self.dispatch(result)
        raise AssertionError("the runner never settled")


@pytest.fixture(scope="module")
def coordinator(tmp_path_factory) -> Coordinator:
    tmp = tmp_path_factory.mktemp("host")
    installed = tmp / "installed" / "rfq-insight"
    shutil.copytree(PLUGIN, installed,
                    ignore=shutil.ignore_patterns("__pycache__"))
    assert (installed / ".claude-plugin" / "plugin.json").is_file()
    folder = tmp / "RFQ with spaces"
    shutil.copytree(ROOT / "examples" / "synthetic-rfq", folder)
    image = Image.new("RGB", (800, 400), "white")
    ImageDraw.Draw(image).text((30, 30), TRANSCRIPTION, fill="black")
    image.save(folder / "review.png")
    return Coordinator(installed / "skills" / "rfq-insight" / "scripts",
                       folder, cwd=tmp)


@pytest.fixture(scope="module")
def finished(coordinator: Coordinator) -> dict:
    """The whole run, remembering what each step returned."""
    log = {}
    # The figures worker is needed even when the user only wants the
    # documents read: `--until figures` dispatches it first ...
    log["first"] = coordinator.next("--until", "figures")
    assert log["first"]["status"] == "dispatch"
    assert log["first"]["stage"] == "figures"
    coordinator.dispatch(log["first"])
    # ... and only reports `stopped` once the stage is complete.
    log["stopped"] = coordinator.next("--until", "figures")
    log["analysis"] = coordinator.next()
    assert log["analysis"]["status"] == "dispatch"
    assert log["analysis"]["stage"] == "analysis"
    coordinator.dispatch(log["analysis"])
    log["done"] = coordinator.loop()
    return log


# ------------------------------------------------------------------ tests
def test_until_figures_dispatches_the_figure_worker_then_stops(finished):
    first = finished["first"]
    (request,) = first["requests"]
    assert request["model"] == "sonnet"
    assert len(request["image_ids"]) == 1
    assert [Path(p).suffix for p in request["outputs"]] == [".txt"]
    assert all(Path(p).parent.name == "image_text" for p in request["outputs"])
    assert finished["stopped"] == {"status": "stopped", "stage": "figures"}


def test_analysis_request_routes_to_opus_with_the_full_instruction_set(
        coordinator, finished):
    (request,) = finished["analysis"]["requests"]
    assert request["model"] == "opus"
    assert "completely" in request["message"]
    names = [Path(p).name for p in request["instructions"]]
    assert names[0].startswith("RFQ_QA_Extraction_Prompt")
    assert "sections.md" in names
    assert all(Path(p).is_relative_to(coordinator.scripts.parent)
               for p in request["instructions"]), "instructions must come " \
        "from the installed copy, not the source repo"
    corpus_txt = coordinator.folder / "output" / ".work" / "corpus.txt"
    assert "19 December 2026" in corpus_txt.read_text(encoding="utf-8"), \
        "the figure transcription must be merged before the analysis reads"


def test_run_finishes_with_a_deck_and_a_twelve_section_report(
        coordinator, finished):
    done = finished["done"]
    assert done["status"] == "done", done
    pptx, report_path = Path(done["pptx"]), Path(done["report"])
    assert pptx.is_file() and pptx.parent == coordinator.folder / "output"
    assert pptx.read_bytes().startswith(b"PK")
    assert pptx.suffix == ".pptx" and report_path.suffix == ".json"
    # Two verification partitions run side by side; the slides are planned
    # by the runner from the answers' markdown, no slide worker is needed.
    assert [s for s, _ in coordinator.served] == [
        "figures", "analysis", "verification", "verification"]
    work = coordinator.folder / "output" / ".work"
    assert json.loads((work / "slide_planning.json").read_text()) == {
        "version": 1, "mode": "deterministic", "parsed": list(range(2, 13))}

    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert len(report["slidePlan"]["questions"]) == 12
    assert [q["number"] for q in report["slidePlan"]["questions"]] == list(range(1, 13))
    assert report["verification"]["status"] == "ok"
    assert report["verification"]["mode"] == "subagent"
    assert report["verification"]["claims"], "the canned answers carry claims"
    assert all(c["verdict"] == "confirmed" for c in report["verification"]["claims"])
    # Every section is scored from its (empty) deductions after the check.
    questions = report["slidePlan"]["questions"]
    assert {(q["confidence"], q["score"]) for q in questions} == {("high", 10)}
    assert all(q["confidenceNotes"].startswith("Confidence 10/10 (High)\n")
               for q in questions)
    timing = next(q for q in questions if q["number"] == 9)
    assert ("Check: 3 statements reviewed by a separate reviewer: "
            "3 confirmed.\n") in timing["confidenceNotes"]
    stages = {s["stage"]: s for s in report["run"]["execution"]["stages"]}
    assert stages["analysis"]["requested_model"] == "opus"
    assert stages["figures"]["requested_model"] == "sonnet"
    assert stages["analysis"]["mode"] == "isolated"
    assert report["run"]["projectName"] == "Cedar Demo"
    assert report["run"]["corpusFiles"] == [SOW, APPENDIX, "review.png"]

    evaluation = (coordinator.folder / "output" / ".work" / "evaluation.txt")
    assert "0 fail" in evaluation.read_text(encoding="utf-8")
    assert json.loads((coordinator.folder / "output" / ".work" /
                       "progress.json").read_text())["pending"] == "done"


def test_opportunity_slide_carries_the_analysis_values(finished):
    report = json.loads(Path(finished["done"]["report"]).read_text(encoding="utf-8"))
    info = _opportunity()
    first = report["slidePlan"]["questions"][0]
    assert first["titleEn"] == "Opportunity Info"
    (slide,) = first["slides"]
    assert slide["layout"] == "opportunity"
    rows = dict(slide["rows"])
    assert rows["Customer"] == info["customer"]
    assert rows["Customer contact"] == info["customer_contact"]
    assert rows["Quotation due date"] == info["quotation_due_date"]
    assert rows["Opportunity name"] == info["opportunity_name"]
    assert rows["Project type"] == info["project_type"]
    assert rows["Lead product"] == info["lead_product"]
    assert rows["Technology"] == info["technology"]
    assert rows["Industry"] == info["industry"]
    # No Salesforce account matches the synthetic customer: the lookup is
    # the runner's, never the model's, and it leaves the placeholders.
    assert rows["Account Name (Site)"] == "Info not available"
    assert rows["KAM (PTE)"] == "Info not available"
    assert slide["products"] == [[info["lead_product"], "100%"]]
    assert slide["skills"] == [[info["skills"][0]["name"], "100%"]]

    from pptx import Presentation
    deck = Presentation(str(finished["done"]["pptx"]))
    texts = [shape.text_frame.text for s in deck.slides for shape in s.shapes
             if shape.has_text_frame]
    tables = [cell.text for s in deck.slides for shape in s.shapes
              if shape.has_table for row in shape.table.rows for cell in row.cells]
    assert any(info["customer"] in t for t in texts + tables)
    assert any(info["project_type"] in t for t in texts + tables)


def test_timing_records_every_dispatched_request(coordinator, finished):
    timing = coordinator.run("timing", coordinator.folder)
    assert timing["status"] == "available"
    by_id = {r["request_id"]: r for r in timing["requests"]}
    assert set(by_id) == {rid for _, rid in coordinator.served}
    for stage, request_id in coordinator.served:
        row = by_id[request_id]
        assert row["stage"] == stage and row["mode"] == "dispatch"
        assert row["started_at"] and row["finished_at"]
        assert row["worker_seconds"] >= 0
        assert row["validation"]["accepted"] is True
    assert {s["stage"] for s in timing["stages"]} >= {
        "extract", "figures", "analysis", "verification", "slides",
        "finalize", "evaluate"}


def test_next_on_a_finished_folder_delivers_again_without_rebuilding(
        coordinator, finished):
    output = coordinator.folder / "output"
    before = sorted(p.name for p in output.glob("*.pptx"))
    assert len(before) == 1
    stat = Path(finished["done"]["pptx"]).stat()
    again = coordinator.next()
    assert again == finished["done"]
    assert sorted(p.name for p in output.glob("*.pptx")) == before
    assert Path(again["pptx"]).stat().st_mtime_ns == stat.st_mtime_ns
    # No worker was needed: `done` only re-ran the quality gate.
    assert len(coordinator.served) == 4


def test_restart_slides_plans_again_without_a_worker(coordinator, finished):
    restarted = coordinator.next("--restart", "slides")
    assert restarted["status"] == "done", restarted
    work = coordinator.folder / "output" / ".work"
    assert list((work / "history").rglob("slide_plan.json")), \
        "the previous plan is archived"
    old, new = Path(finished["done"]["pptx"]), Path(restarted["pptx"])
    assert new != old and new.is_file() and old.is_file(), \
        "an existing presentation is never overwritten"
    assert new.stem == old.stem + "_2"
    assert len(coordinator.served) == 4, "no worker was needed"


def test_an_unparsable_section_sends_the_deck_to_the_slide_worker(
        coordinator, finished):
    coordinator.answers = {4: UNPARSABLE_4}
    restarted = coordinator.next("--restart", "analysis")
    assert restarted["status"] == "dispatch" and restarted["stage"] == "analysis"
    coordinator.dispatch(restarted)
    done = coordinator.loop()
    assert done["status"] == "done", done
    work = coordinator.folder / "output" / ".work"
    assert json.loads((work / "slide_planning.json").read_text()) == {
        "version": 1, "mode": "worker", "unparsable": [4]}
    (request_stage, _id) = coordinator.served[-1]
    assert request_stage == "slides"
    slides = [r for r in coordinator.served if r[0] == "slides"]
    assert len(slides) == 1
    assert Path(done["pptx"]).stem == Path(finished["done"]["pptx"]).stem + "_3"
    report = json.loads(Path(done["report"]).read_text(encoding="utf-8"))
    assert len(report["slidePlan"]["questions"]) == 12
