import importlib.util
import os
import re
import zipfile

from pptx import Presentation
from pptx.util import Inches

from conftest import REPO_ROOT, SKILL_ROOT

TEMPLATE = os.path.join(SKILL_ROOT, "assets", "avl_template.pptx")


def _load_builder():
    path = os.path.join(REPO_ROOT, "tools", "make_template.py")
    spec = importlib.util.spec_from_file_location("make_template", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_template_is_a_master_without_slides():
    prs = Presentation(TEMPLATE)
    assert len(prs.slides) == 0
    assert abs(prs.slide_width / 914400 - 13.333) < 0.01
    assert abs(prs.slide_height / 914400 - 7.5) < 0.01
    names = [l.name for l in prs.slide_layouts]
    assert "Only Title" in names
    assert any(n.startswith("1_Title_") for n in names)
    assert any(n.startswith("Final Slide") for n in names)


def test_template_carries_no_slide_parts_or_large_media():
    with zipfile.ZipFile(TEMPLATE) as z:
        names = z.namelist()
        assert not [n for n in names if n.startswith("ppt/slides/")]
        assert all(z.getinfo(n).file_size <= 1_000_000
                   for n in names if n.startswith("ppt/media/"))


def test_template_leaks_nothing_from_the_source_slides():
    with zipfile.ZipFile(TEMPLATE) as z:
        names = z.namelist()
        assert names[0] == "[Content_Types].xml"
        assert not [n for n in names if n.startswith(("ppt/notesSlides/", "[trash]/",
                                                       "ppt/changesInfos/"))]
        assert "docProps/thumbnail.jpeg" not in names
        text = {n: z.read(n).decode("utf-8", "ignore") for n in names
                if n.endswith((".xml", ".rels"))}
    assert not [n for n, t in text.items() if re.search(r"slides/slide\d+\.xml", t)]
    assert not [n for n, t in text.items() if "notesSlide" in t or "thumbnail" in t]
    assert "<p14:sectionLst" not in text["ppt/presentation.xml"]
    assert "Folientitel" not in text["docProps/app.xml"]
    assert "<Slides>0</Slides>" in text["docProps/app.xml"]


def test_every_media_part_is_referenced_by_a_remaining_part():
    with zipfile.ZipFile(TEMPLATE) as z:
        media = {n for n in z.namelist() if n.startswith("ppt/media/")}
        referenced = set()
        for n in z.namelist():
            if n.endswith(".rels"):
                for target in re.findall(r'Target="\.\./media/([^"]+)"',
                                         z.read(n).decode("utf-8")):
                    referenced.add("ppt/media/" + target)
    assert media and media == referenced


def test_strip_to_master_drops_slides_notes_and_orphan_media(tmp_path):
    builder = _load_builder()
    picture = tmp_path / "pic.png"
    # smallest valid PNG: 1x1 pixel
    picture.write_bytes(bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000d49444154789c6360f8cfc000000301010018dd8db00000000049454e44ae426082"))
    src = tmp_path / "src.pptx"
    prs = Presentation()
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Customer secret %d" % i
        slide.shapes.add_picture(str(picture), Inches(1), Inches(1))
        slide.notes_slide.notes_text_frame.text = "private note"
    prs.save(src)
    with zipfile.ZipFile(src) as z:
        assert any(n.startswith("ppt/notesSlides/") for n in z.namelist())
        assert any(n.startswith("ppt/media/") for n in z.namelist())

    out = tmp_path / "out.pptx"
    result = builder.strip_to_master(str(src), str(out))

    assert result["slides_removed"] == 3
    assert result["media_removed"] >= 1
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert names[0] == "[Content_Types].xml"
        assert not [n for n in names if n.startswith(("ppt/slides/", "ppt/notesSlides/"))]
        assert not [n for n in names if n.startswith("ppt/media/")]
        blob = b"".join(z.read(n) for n in names)
        assert b"Customer secret" not in blob and b"private note" not in blob
        assert b"/ppt/slides/" not in z.read("[Content_Types].xml")
    reopened = Presentation(out)
    assert len(reopened.slides) == 0
    assert len(reopened.slide_layouts) == len(prs.slide_layouts)


def test_template_carries_no_document_metadata():
    with zipfile.ZipFile(TEMPLATE) as z:
        names = z.namelist()
        assert "docProps/custom.xml" not in names
        assert not [n for n in names if n.startswith("customXml/")]
        core = z.read("docProps/core.xml").decode("utf-8")
        text = "".join(z.read(n).decode("utf-8", "ignore") for n in names
                       if n.endswith((".xml", ".rels")))
    assert "<dc:creator>RFQ Insight</dc:creator>" in core
    assert "<cp:lastModifiedBy>RFQ Insight</cp:lastModifiedBy>" in core
    assert "customXml" not in text and "custom-properties" not in text


def test_template_has_no_author_or_footer_text_boxes():
    with zipfile.ZipFile(TEMPLATE) as z:
        for n in z.namelist():
            if not re.fullmatch(r"ppt/(slideMasters/slideMaster|slideLayouts/slideLayout)\d+\.xml", n):
                continue
            for shape in re.findall(r"<p:sp(?:\s[^>]*)?>.*?</p:sp>",
                                   z.read(n).decode("utf-8"), re.S):
                texts = re.findall(r"<a:t>(.*?)</a:t>", shape, re.S)
                if "<p:ph" not in shape:
                    assert not any("|" in t for t in texts), n
                if 'name="txt_Autor' in shape:
                    assert not any(t.strip() for t in texts), n


def _box(shape_id, name, text, placeholder=False):
    ph = "<p:nvPr><p:ph type=\"ftr\"/></p:nvPr>" if placeholder else "<p:nvPr/>"
    return (
        '<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr txBox="1"/>%s</p:nvSpPr>'
        '<p:spPr/><p:txBody><a:bodyPr/><a:p><a:r><a:rPr lang="en-GB"/><a:t>%s</a:t></a:r></a:p>'
        "</p:txBody></p:sp>" % (shape_id, name, ph, text))


def _deck_with_named_boxes(tmp_path, layout_box, creator="Some Author"):
    """A deck whose master/layout carry author text, as the AVL source does."""
    prs = Presentation()
    prs.core_properties.author = creator
    prs.core_properties.last_modified_by = creator
    src = tmp_path / "named.pptx"
    prs.save(src)
    edited = tmp_path / "named_edited.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(edited, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/slideMasters/slideMaster1.xml":
                data = data.decode("utf-8").replace("</p:spTree>", _box(
                    90, "footer_ProjectInfo", "Some Author | DEPT | 1 Januar 2030 |"
                ) + "</p:spTree>").encode("utf-8")
            if item.filename == "ppt/slideLayouts/slideLayout1.xml":
                data = data.decode("utf-8").replace(
                    "</p:spTree>", layout_box + "</p:spTree>").encode("utf-8")
            zout.writestr(item, data)
    return edited


def _all_xml(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n).decode("utf-8", "ignore") for n in z.namelist()
                if n.endswith((".xml", ".rels"))}


def test_strip_to_master_scrubs_author_text_generically(tmp_path):
    builder = _load_builder()
    table = (
        '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="92" name="tbl_Info"/>'
        "</p:nvGraphicFramePr><a:tbl><a:tr><a:tc><a:txBody><a:p><a:r><a:t>Author(s):"
        "</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r>"
        "<a:t>Some Author</a:t></a:r></a:p></a:txBody></a:tc></a:tr></a:tbl>"
        "</p:graphicFrame>")
    src = _deck_with_named_boxes(tmp_path, _box(91, "txt_Autor1", "Some Author") + table)
    out = tmp_path / "out.pptx"
    builder.strip_to_master(str(src), str(out))
    xml = _all_xml(out)
    assert "Author(s):" in xml["ppt/slideLayouts/slideLayout1.xml"]
    for name, text in xml.items():
        assert "Some Author" not in text, name
        assert "DEPT" not in text, name
    core = xml["docProps/core.xml"]
    assert "<dc:creator>RFQ Insight</dc:creator>" in core
    # the boxes remain (only their text is blanked) so layouts keep their geometry
    assert 'name="footer_ProjectInfo"' in xml["ppt/slideMasters/slideMaster1.xml"]
    assert 'name="txt_Autor1"' in xml["ppt/slideLayouts/slideLayout1.xml"]


def test_strip_to_master_keeps_placeholder_prompts(tmp_path):
    builder = _load_builder()
    src = _deck_with_named_boxes(tmp_path, _box(91, "Footer Placeholder", "Page footer", True))
    out = tmp_path / "out.pptx"
    builder.strip_to_master(str(src), str(out))
    assert "Page footer" in _all_xml(out)["ppt/slideLayouts/slideLayout1.xml"]


def test_strip_to_master_fails_when_an_author_name_survives(tmp_path):
    builder = _load_builder()
    src = _deck_with_named_boxes(tmp_path, _box(91, "Some label", "Some Author"))
    try:
        builder.strip_to_master(str(src), str(tmp_path / "out.pptx"))
    except ValueError as exc:
        assert "ppt/slideLayouts/slideLayout1.xml" in str(exc)
    else:
        raise AssertionError("expected ValueError")
