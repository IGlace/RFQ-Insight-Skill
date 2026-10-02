"""Folder walk, per-file extraction, and the corpus checkpoint."""

import io
import json
import os

import extract_corpus as ec


def _docx(path, paragraph="The system shall deliver 450 Nm."):
    import docx
    d = docx.Document()
    d.add_paragraph(paragraph)
    d.save(path)


def _noise_png(w, h):
    """A PNG big enough in both bytes and pixels to survive the filter."""
    from PIL import Image
    buf = io.BytesIO()
    Image.frombytes("RGB", (w, h), os.urandom(w * h * 3)).save(buf, format="PNG")
    return buf.getvalue()


def _docx_with_picture(path, png_bytes, paragraph="Body text."):
    import docx
    from docx.shared import Inches
    png = os.path.join(os.path.dirname(path), "_tmp_pic.png")
    with open(png, "wb") as fh:
        fh.write(png_bytes)
    d = docx.Document()
    d.add_paragraph(paragraph)
    d.add_picture(png, width=Inches(2))
    d.save(path)
    os.remove(png)


def _rfq_folder(tmp_path):
    folder = tmp_path / "Acme V6"
    (folder / "annexes").mkdir(parents=True)
    _docx(str(folder / "SOR.docx"))
    (folder / "notes.txt").write_text("Quotation due 20.06.2026",
                                      encoding="utf-8")
    _docx(str(folder / "annexes" / "Annex A.docx"), "Annex content.")
    (folder / "broken.docx").write_bytes(b"not a docx")
    (folder / "onenote.one").write_bytes(b"\x00\x01")
    return str(folder)


def test_walk_finds_supported_files_recursively_and_sorted(tmp_path):
    folder = _rfq_folder(tmp_path)
    supported, ignored = ec.walk_supported(folder)
    names = [os.path.basename(p) for p in supported]
    assert names == ["Annex A.docx", "SOR.docx", "broken.docx", "notes.txt"]
    assert ignored == ["onenote.one"]


def test_walk_never_descends_into_our_own_output_folder(tmp_path):
    folder = _rfq_folder(tmp_path)
    prior = os.path.join(folder, "output")
    os.makedirs(prior)
    _docx(os.path.join(prior, "Previous_RFQ_Analysis.docx"))
    supported, _ = ec.walk_supported(folder)
    assert all("output" not in os.path.relpath(p, folder).split(os.sep)
               for p in supported)


def test_build_corpus_carries_the_per_file_contract(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = ec.build_corpus(folder, "Acme V6")

    assert corpus["projectName"] == "Acme V6"
    assert corpus["oppId"] == ""
    by_name = {f["filename"]: f for f in corpus["filesContents"]}
    assert set(by_name) == {"SOR.docx", "notes.txt", "Annex A.docx"}

    sor = by_name["SOR.docx"]
    assert sor["extension"] == "docx"
    assert sor["category"] == "office_doc"
    assert sor["extraction_method"] == "python-docx"
    assert sor["path"] == "SOR.docx"
    assert sor["file_size"] > 0
    assert "450 Nm" in sor["extracted_text"]
    assert by_name["Annex A.docx"]["path"] == os.path.join("annexes",
                                                           "Annex A.docx")


def test_unreadable_file_is_recorded_with_a_reason_not_dropped(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = ec.build_corpus(folder, "Acme V6")
    skipped = corpus["extraction"]["skipped"]
    assert [s["filename"] for s in skipped] == ["broken.docx"]
    assert "corrupt" in skipped[0]["reason"]
    assert corpus["extraction"]["ignored"] == ["onenote.one"]
    assert corpus["extraction"]["parsed"] == 3


def test_preprocess_stats_are_recorded_per_file(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, _ = ec.build_corpus(folder, "Acme V6")
    stats = corpus["extraction"]["preprocess"]
    assert set(stats) == {"SOR.docx", "notes.txt", "Annex A.docx"}
    assert stats["SOR.docx"]["chars_after"] > 0


def test_cli_writes_the_three_checkpoints(tmp_path):
    folder = _rfq_folder(tmp_path)
    assert ec.main([folder, "--project", "Acme V6"]) == 0

    work = ec.work_dir(folder)
    with open(os.path.join(work, "corpus.json"), encoding="utf-8") as fh:
        corpus = json.load(fh)
    with open(os.path.join(work, "images.json"), encoding="utf-8") as fh:
        manifest = json.load(fh)

    assert corpus["projectName"] == "Acme V6"
    assert os.path.isdir(os.path.join(work, "images"))
    assert manifest["images"] == []          # nothing embedded in the fixtures


def test_project_name_defaults_to_the_folder_name(tmp_path):
    folder = _rfq_folder(tmp_path)
    assert ec.main([folder]) == 0
    with open(os.path.join(ec.work_dir(folder), "corpus.json"),
              encoding="utf-8") as fh:
        assert json.load(fh)["projectName"] == "Acme V6"


def test_empty_folder_exits_one(tmp_path):
    folder = tmp_path / "Empty"
    folder.mkdir()
    assert ec.main([str(folder)]) == 1


def test_summary_is_plain_language(tmp_path):
    folder = _rfq_folder(tmp_path)
    corpus, raw = ec.build_corpus(folder, "Acme V6")
    from rfq.images import filter_images, write_images
    kept, dropped = filter_images(raw)
    manifest = write_images(kept, dropped,
                            os.path.join(ec.work_dir(folder), "images"))
    text = "\n".join(ec.summary_lines(corpus, manifest))

    assert "Read 3 of 4 files" in text
    assert "broken.docx" in text and "corrupt" in text
    for jargon in ("corpus", "extraction_method", "Traceback", "None"):
        assert jargon not in text


# ------------------------------------------- image identity across the corpus
def _colliding_basenames(tmp_path):
    """annexes/Spec.docx and attachments/Spec.docx, each with its OWN
    picture. The reviewer's probe: two documents, one basename."""
    folder = tmp_path / "Acme V6"
    (folder / "annexes").mkdir(parents=True)
    (folder / "attachments").mkdir(parents=True)
    annex_png, attach_png = _noise_png(300, 300), _noise_png(301, 301)
    _docx_with_picture(str(folder / "annexes" / "Spec.docx"), annex_png,
                       "Annex body.")
    _docx_with_picture(str(folder / "attachments" / "Spec.docx"), attach_png,
                       "Attachment body.")
    return str(folder), annex_png, attach_png


def test_same_named_documents_get_distinct_image_ids(tmp_path):
    """An id built from the basename alone gave both documents
    'Spec.docx__img01' - so one picture's bytes overwrote the other's and
    the merge handed a document a chart it does not contain."""
    folder, _, _ = _colliding_basenames(tmp_path)
    _, raw = ec.build_corpus(folder, "Acme V6")

    ids = [i.image_id for i in raw]
    assert len(ids) == 2
    assert len(set(ids)) == 2, f"colliding image ids: {ids}"
    assert sorted(ids) == ["annexes_Spec.docx__img01",
                           "attachments_Spec.docx__img01"]
    # `source` stays the display name - it is what the user is shown.
    assert {i.source for i in raw} == {"Spec.docx"}


def test_same_named_documents_keep_their_own_picture_on_disk(tmp_path):
    """The end of the probe: two ids, two files, two different images."""
    folder, annex_png, attach_png = _colliding_basenames(tmp_path)
    assert ec.main([folder, "--project", "Acme V6"]) == 0

    work = ec.work_dir(folder)
    with open(os.path.join(work, "images.json"), encoding="utf-8") as fh:
        manifest = json.load(fh)

    ids = [e["id"] for e in manifest["images"]]
    assert sorted(ids) == ["annexes_Spec.docx__img01",
                           "attachments_Spec.docx__img01"]
    files = sorted(os.listdir(os.path.join(work, "images")))
    assert files == ["annexes_Spec.docx__img01.png",
                     "attachments_Spec.docx__img01.png"]

    on_disk = {}
    for name in files:
        with open(os.path.join(work, "images", name), "rb") as fh:
            on_disk[name] = fh.read()
    assert on_disk["annexes_Spec.docx__img01.png"] == annex_png
    assert on_disk["attachments_Spec.docx__img01.png"] == attach_png


def test_each_same_named_document_carries_only_its_own_marker(tmp_path):
    """The other half of the collision: identical marker text in both
    documents is what let the merge resolve the wrong one."""
    from rfq.extractors import marker_for
    folder, _, _ = _colliding_basenames(tmp_path)
    assert ec.main([folder, "--project", "Acme V6"]) == 0

    with open(os.path.join(ec.work_dir(folder), "corpus.json"),
              encoding="utf-8") as fh:
        by_path = {f["path"]: f["extracted_text"]
                   for f in json.load(fh)["filesContents"]}

    annex = by_path[os.path.join("annexes", "Spec.docx")]
    attach = by_path[os.path.join("attachments", "Spec.docx")]
    assert marker_for("annexes_Spec.docx__img01") in annex
    assert marker_for("annexes_Spec.docx__img01") not in attach
    assert marker_for("attachments_Spec.docx__img01") in attach
    assert marker_for("attachments_Spec.docx__img01") not in annex


def test_identical_copies_of_a_document_still_get_their_own_id(tmp_path):
    """The commoner case: the same file filed twice. Same bytes, so the
    second picture is a genuine duplicate - but it must be recorded under
    its OWN id, not silently merged into the first document's."""
    folder = tmp_path / "Acme V6"
    (folder / "annexes").mkdir(parents=True)
    (folder / "attachments").mkdir(parents=True)
    png = _noise_png(300, 300)
    _docx_with_picture(str(folder / "annexes" / "Spec.docx"), png, "Body.")
    _docx_with_picture(str(folder / "attachments" / "Spec.docx"), png, "Body.")

    _, raw = ec.build_corpus(str(folder), "Acme V6")
    assert sorted(i.image_id for i in raw) == ["annexes_Spec.docx__img01",
                                               "attachments_Spec.docx__img01"]


def test_a_top_level_file_keeps_its_plain_name_in_the_image_id(tmp_path):
    """Readability: the id reaches the user, so the usual case stays the
    filename they recognise."""
    folder = tmp_path / "Acme V6"
    folder.mkdir()
    _docx_with_picture(str(folder / "SOR.docx"), _noise_png(300, 300))
    _, raw = ec.build_corpus(str(folder), "Acme V6")
    assert [i.image_id for i in raw] == ["SOR.docx__img01"]


def test_two_paths_that_slug_alike_are_still_told_apart(tmp_path):
    """'a b/S.docx' and 'a_b/S.docx' slug to the same string. The counter
    is what makes the id unique rather than merely unlikely to collide."""
    folder = tmp_path / "Acme V6"
    (folder / "a b").mkdir(parents=True)
    (folder / "a_b").mkdir(parents=True)
    _docx_with_picture(str(folder / "a b" / "S.docx"), _noise_png(300, 300))
    _docx_with_picture(str(folder / "a_b" / "S.docx"), _noise_png(301, 301))

    _, raw = ec.build_corpus(str(folder), "Acme V6")
    ids = [i.image_id for i in raw]
    assert len(set(ids)) == 2, f"colliding image ids: {ids}"


# ------------------------------------------ no marker without a picture behind
def test_a_filtered_out_picture_leaves_no_transcription_marker(tmp_path):
    """Every embedded picture gets a marker on the way out of the parser,
    but most are filtered away and will never be transcribed. A marker with
    nothing coming makes the analysis hedge answers over content that does
    not exist."""
    from rfq.extractors import marker_for
    folder = tmp_path / "Acme V6"
    folder.mkdir()
    _docx_with_picture(str(folder / "SOR.docx"), _noise_png(16, 16),
                       "Requirement text.")

    assert ec.main([str(folder), "--project", "Acme V6"]) == 0
    work = ec.work_dir(str(folder))
    with open(os.path.join(work, "images.json"), encoding="utf-8") as fh:
        manifest = json.load(fh)
    assert manifest["images"] == []
    assert [d["reason"] for d in manifest["filtered"]] == ["too_small"]

    with open(os.path.join(work, "corpus.json"), encoding="utf-8") as fh:
        corpus = json.load(fh)
    text = corpus["filesContents"][0]["extracted_text"]
    assert "transcription pending" not in text
    assert marker_for("SOR.docx__img01") not in text
    assert "Requirement text." in text
    assert corpus["extraction"]["chars"] == len(text)


def test_a_kept_picture_keeps_its_marker(tmp_path):
    """The other side of the same rule: drop_markers must not take the
    placeholder for a picture that IS going to be transcribed."""
    from rfq.extractors import marker_for
    folder = tmp_path / "Acme V6"
    folder.mkdir()
    _docx_with_picture(str(folder / "SOR.docx"), _noise_png(300, 300))

    assert ec.main([str(folder), "--project", "Acme V6"]) == 0
    with open(os.path.join(ec.work_dir(str(folder)), "corpus.json"),
              encoding="utf-8") as fh:
        text = json.load(fh)["filesContents"][0]["extracted_text"]
    assert marker_for("SOR.docx__img01") in text


def test_dropping_a_marker_leaves_the_surrounding_text_intact(tmp_path):
    from rfq.extractors import marker_for
    marker = marker_for("SOR.docx__img01")
    corpus = {"filesContents": [
        {"filename": "SOR.docx",
         "extracted_text": f"Before.\n\n{marker}\n\nAfter."}],
        "extraction": {"chars": 0}}
    assert ec.drop_markers(corpus, [{"id": "SOR.docx__img01"}]) == 1
    assert corpus["filesContents"][0]["extracted_text"] == "Before.\n\nAfter."


def test_the_skipped_picture_summary_names_the_real_reason(tmp_path):
    """'logos or page furniture' was printed for pictures that were over
    the per-run limit, or not readable at all. Both are untrue."""
    filtered = [
        {"id": "a", "reason": "too_small", "explanation": "x"},
        {"id": "b", "reason": "cap_reached", "explanation": "y"},
        {"id": "c", "reason": "unreadable", "explanation": "z"},
    ]
    lines = ec.filtered_lines(filtered)
    text = "\n".join(lines)
    assert len(lines) == 3
    assert "1 picture(s) that are too small" in text
    assert "beyond the limit" in text
    assert "could not open" in text
    for jargon in ("cap_reached", "too_small", "unreadable"):
        assert jargon not in text
