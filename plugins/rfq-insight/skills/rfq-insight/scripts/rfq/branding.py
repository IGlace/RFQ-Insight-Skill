"""The AVL slide master and the deck's footer identity.

Every deck is built on assets/avl_template.pptx (the AVL master with its
layouts and no slides). The master already draws the footer: "Internal" on
the left, the live slide number, the logo and the gradient bar. The one
thing the renderer writes into it is the right-hand box
(`footer_ProjectInfo`): "<author> | <department> | <D Month YYYY> |".

Author and department are not in the RFQ, so they are configured once per
machine. Resolution order: the explicit argument (`rfq_run.py next
--author/--department`, saved with the run), the RFQ_INSIGHT_AUTHOR /
RFQ_INSIGHT_DEPARTMENT environment variables, the `author` / `department`
keys of <toolbox home>/config.json, then the operating-system account name
and "RFQ Insight" (recorded as a fallback in the run report).
"""

import datetime
import getpass
import json
import os
from collections.abc import Mapping

TEMPLATE_PATH = os.path.abspath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), os.pardir, os.pardir,
    "assets", "avl_template.pptx"))

LAYOUT_CONTENT = "Only Title"
LAYOUT_COVER = "1_Title_ AVL Company"
LAYOUT_END = "Final Slide with Backgroundimage English"

FOOTER_BOX = "footer_ProjectInfo"
FALLBACK_DEPARTMENT = "RFQ Insight"
# English month names, independent of the machine's locale (%B is not).
MONTHS = ("January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December")


def load_template():
    """A fresh Presentation on the AVL master (python-pptx is imported here
    so that resolving the identity never needs it)."""
    from pptx import Presentation
    return Presentation(TEMPLATE_PATH)


def long_date(iso: str) -> str:
    """'2026-10-23' -> '23 October 2026'; anything else comes back stripped."""
    value = str(iso or "").strip()
    try:
        day = datetime.date.fromisoformat(value[:10])
    except ValueError:
        return value
    return f"{day.day} {MONTHS[day.month - 1]} {day.year}"


def _home_dir() -> str:
    # bootstrap.HOME_DIR, without importing bootstrap (it has side effects).
    return (os.environ.get("RFQ_INSIGHT_HOME")
            or os.path.join(os.path.expanduser("~"), ".rfq-insight"))


def _config(home: str) -> dict:
    try:
        with open(os.path.join(home, "config.json"), encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def _clean(value) -> str:
    return value.strip() if isinstance(value, str) else ""


def _os_user() -> str:
    try:
        return getpass.getuser() or "RFQ Insight"
    except Exception:          # no login name (some containers / services)
        return "RFQ Insight"


def resolve_identity(author: str | None = None, department: str | None = None,
                     env: Mapping = os.environ, home: str | None = None) -> dict:
    """{"author", "department", "fallback"}: argument, then environment,
    then config.json, then the OS account name / "RFQ Insight"."""
    config = _config(home if home is not None else _home_dir())
    fallback = False
    found = {}
    for key, given, variable, default in (
            ("author", author, "RFQ_INSIGHT_AUTHOR", _os_user),
            ("department", department, "RFQ_INSIGHT_DEPARTMENT",
             lambda: FALLBACK_DEPARTMENT)):
        value = (_clean(given) or _clean(env.get(variable))
                 or _clean(config.get(key)))
        if not value:
            value, fallback = default(), True
        found[key] = value
    return {"author": found["author"], "department": found["department"],
            "fallback": fallback}


def footer_text(author: str, department: str, iso_date: str) -> str:
    return f"{author} | {department} | {long_date(iso_date)} |"


def write_text(shape, value: str) -> None:
    """Set a template text box's text in the first run of its first
    paragraph, so the run's formatting (font, size, colour) is kept; any
    further runs are emptied."""
    paragraph = shape.text_frame.paragraphs[0]
    runs = paragraph.runs
    run = runs[0] if runs else paragraph.add_run()
    run.text = value
    for extra in runs[1:]:
        extra.text = ""


def set_footer(prs, author: str, department: str, iso_date: str) -> None:
    """Write the footer identity into the master's `footer_ProjectInfo` box
    (Verdana 10 pt, kept). The security label and slide-number boxes are
    not touched."""
    master = prs.slide_masters[0]
    box = next((s for s in master.shapes if s.name == FOOTER_BOX), None)
    if box is None or not box.has_text_frame:
        raise ValueError(f"the slide master has no {FOOTER_BOX} text box")
    write_text(box, footer_text(author, department, iso_date))
