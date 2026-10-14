"""RFQ Insight deck renderer — pinned, deterministic (production).

Styled after the client's RFQ Insight template (AVL blue->cyan->green->lime
gradient, Verdana, compact one-topic-per-slide flow with a Table of Contents
instead of per-question section covers), with DIAGRAM layouts for the
sections the client flagged as picture-worthy:

  * "gantt"      — ONE-slide Gantt: year/quarter calendar grid, project (or
                   per-phase) bars and every milestone at its proportional
                   time position (the sales-preferred Timing picture).
  * "timeline"   — milestone graphic: numbered dots on a gradient axis, the
                   milestone name/date/desc in the card (labels never sit in
                   the circles, so they cannot truncate). Retained as the
                   fallback when milestone dates cannot be parsed.
  * "phases"     — chevron process diagram (Development Phases).
  * "collab"     — hub-and-spoke collaboration model (Collaboration Approach).
  * "ownership"  — Scope-of-Work matrix with color-coded owner cells + legend.
  * "pictures"   — up to four product pictures copied from the RFQ, each
                   with caption and source (built by rfq/delivery.py, never by
                   the slide-plan model).

Everything else (bullets / table / twoColumn / prosCons / kpi / prose) is
schema-compatible with rfq_insight.schema.SlidePlan. Self-contained: no
network, gradients generated with Pillow, AVL logo from the local assets
module.
"""

import io
import copy
import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from PIL import Image

from .assets import logo_white_stream  # embedded AVL logo bytes
from .dates import display_date
from .figures import deck_picture, resolve_picture_file
from .schema import ADDITIONAL_INFO_ID, SOURCES

# ------------------------------------------------------------- brand (RFQ Insight)
# Theme colors extracted from the client's RFQ Insight example deck.
BLUE      = "005A99"   # accent1 — primary AVL blue
CYAN      = "00A5C8"   # accent2
GREEN     = "13A538"   # accent3
LIME      = "AFCB08"   # accent4
GRAY      = "888888"   # accent6
GRAY_SOFT = "BFBFBF"   # accent5
INK       = "1A1A1A"
INK_SOFT  = "5A5F5C"
LIGHT     = "F2F6F9"   # cool light fill (blue-tinted)
LIGHT2    = "E8EEF3"
WHITE     = "FFFFFF"
LINE      = "D7DEE4"
WARN      = "C0392B"
AMBER     = "E9B300"

HF = "Verdana"   # the template's font, headers
BF = "Verdana"   # body

PAGE_W, PAGE_H, M = 13.333, 7.5, 0.6
FOOT_H = 0.34                      # gradient footer bar height
LOGO_AR = 480 / 206.0

# Source band (F2 traceability): one gray 8 pt line naming where the slide's
# information comes from, just above the footer bar. CONTENT_BOTTOM is the
# lowest y ANY content shape may reach - every layout budgets against it, so
# nothing can run into the source line; SOURCE_GAP keeps a visible gap
# between the lowest content and the line.
SOURCE_Y = PAGE_H - FOOT_H - 0.22
SOURCE_H = 0.18
SOURCE_GAP = 0.05
CONTENT_BOTTOM = SOURCE_Y - SOURCE_GAP

# Owner color coding for the Scope-of-Work matrix (+ legend). "customer" also
# matches the customer abbreviation supplied in the plan (e.g. "VF").
#
# "joint" is NOT a fourth owner: it marks a task the RFQ assigns to two parties
# at once. Without it a blank cell meant two different things - "the RFQ
# assigns nobody" and "the RFQ assigns two parties" - rendered identically, so
# the matrix lost a fact it was holding. Amber rather than a party color on
# purpose: the cell reads as a flag to settle, not as an answer, which keeps
# the client's confirmed single-owner format from drifting back into a RASIC
# matrix. The parties themselves are named in the Risk Detection entry.
OWNER_COLORS = {"avl": GREEN, "customer": BLUE, "supplier": GRAY,
                "joint": AMBER}


def C(h):
    return RGBColor.from_string(h)


def _hx(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


# ------------------------------------------------------------------- gradients
def _gradient_png(stops, w=1920, h=1080, vertical=False):
    """Multi-stop left->right (or top->bottom) gradient PNG."""
    img = Image.new("RGB", (w, h))
    px = img.load()
    n = len(stops) - 1
    for y in range(h):
        for x in range(w):
            t = (y / (h - 1)) if vertical else (x / (w - 1))
            seg = min(int(t * n), n - 1)
            lt = t * n - seg
            c0, c1 = stops[seg], stops[seg + 1]
            px[x, y] = tuple(int(c0[i] + (c1[i] - c0[i]) * lt) for i in range(3))
    buf = io.BytesIO(); img.save(buf, "PNG"); buf.seek(0)
    return buf


BRAND_STOPS = [_hx(BLUE), _hx(CYAN), _hx(GREEN), _hx(LIME)]
GRAD_COVER = lambda: _gradient_png(BRAND_STOPS)                       # noqa: E731
GRAD_BAR = lambda: _gradient_png(BRAND_STOPS, w=1600, h=40)           # noqa: E731


# ------------------------------------------------------------------ primitives
def slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg_solid(s, color):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = C(color)


def bg_image(s, buf):
    s.shapes.add_picture(buf, 0, 0, Inches(PAGE_W), Inches(PAGE_H))


def pic(s, buf, x, y, w, h):
    s.shapes.add_picture(buf, Inches(x), Inches(y), Inches(w), Inches(h))


def rect(s, x, y, w, h, fill=None, line=None, line_w=1.0, rounded=False, radius=0.08):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = C(line); shp.line.width = Pt(line_w)
    return shp


def shape(s, kind, x, y, w, h, fill, line=None, line_w=1.5):
    shp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = C(line); shp.line.width = Pt(line_w)
    return shp


def connector(s, x0, y0, x1, y1, color=GRAY, weight=1.5, dashed=False):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    cn.shadow.inherit = False
    cn.line.color.rgb = C(color); cn.line.width = Pt(weight)
    if dashed:
        from pptx.oxml.ns import qn
        ln = cn.line._get_or_add_ln()
        d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(d)
    return cn


def text(s, runs, x, y, w, h, size=12, color=INK, bold=False, italic=False,
         font=BF, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True,
         space_after=0, line_spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    items = runs if isinstance(runs, list) else [runs]
    for i, ln in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic
        f.name = font; f.color.rgb = C(color)
    return tb


def bullets(s, items, x, y, w, h, size=13, color=INK, gap=8, marker_color=BLUE):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        rb = p.add_run(); rb.text = "▪  "
        rb.font.size = Pt(size); rb.font.name = BF
        rb.font.color.rgb = C(marker_color); rb.font.bold = True
        rt = p.add_run(); rt.text = it
        rt.font.size = Pt(size); rt.font.name = BF; rt.font.color.rgb = C(color)
    return tb


# ---------------------------------------------------------------------- chrome
def _grad_underbar(s, x, y, w, h=0.05):
    s.shapes.add_picture(GRAD_BAR(), Inches(x), Inches(y), Inches(w), Inches(h))


def title(s, t):
    # Title box always stops before the confidence-chip zone, and the font
    # auto-shrinks so long titles stay on one line instead of colliding.
    avail = PAGE_W - 2 * M - TITLE_RIGHT_GUARD
    size = 26
    for cand in (26, 22, 19, 16):
        size = cand
        if len(t) * 0.0205 * cand <= avail:   # ~Verdana-bold width estimate
            break
    text(s, t, M, 0.42, avail, 0.62, size=size, bold=True,
         color=INK, font=HF, anchor=MSO_ANCHOR.MIDDLE)
    _grad_underbar(s, M, 1.12, 2.2)


def footer(s, page_label, project, date):
    s.shapes.add_picture(GRAD_BAR(), 0, Inches(PAGE_H - FOOT_H),
                         Inches(PAGE_W), Inches(FOOT_H))
    yy = PAGE_H - FOOT_H + 0.05
    text(s, f"Internal   ·   {page_label}", M, yy, 4, 0.24, size=8.5, color=WHITE)
    proj = " – ".join(p.strip() for p in str(project).split("\n") if p.strip())
    if len(proj) > 70:
        proj = proj[:67] + "…"
    text(s, f"RFQ Insight  |  {proj}  |  {date}", PAGE_W - M - 6.6, yy, 5.3,
         0.24, size=8.5, color=WHITE, align=PP_ALIGN.RIGHT, wrap=False)
    w = 0.62
    pic(s, logo_white_stream(), PAGE_W - M - w + 0.55, yy - 0.02, w, w / LOGO_AR)


CONF_COLORS = {"high": GREEN, "medium": AMBER, "low": WARN}

CHIP_W, CHIP_H = 1.85, 0.32          # confidence chip, pinned to the top-right corner
TITLE_RIGHT_GUARD = CHIP_W + 0.35    # title text may never enter this zone


def _confidence_chip(s, conf):
    """Colored chip pinned to the TOP-RIGHT corner showing the section's confidence."""
    conf = str(conf or "").strip().lower()
    color = CONF_COLORS.get(conf)
    if not color:
        return
    x = PAGE_W - M - CHIP_W
    y = 0.45
    rect(s, x, y, CHIP_W, CHIP_H, fill=color, rounded=True, radius=0.5)
    text(s, f"CONFIDENCE: {conf.upper()}", x, y, CHIP_W, CHIP_H, size=9, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _source_line(s, src):
    """The slide's source line. format_source_line() caps it at 150
    characters, which fits the content width on one line at 8 pt."""
    if not src:
        return
    text(s, src, M, SOURCE_Y, PAGE_W - 2 * M, SOURCE_H, size=8, color=GRAY,
         wrap=False, anchor=MSO_ANCHOR.MIDDLE)


def content_slide(prs, ctx, heading, chip=True):
    """White content slide with title and footer. Every physical page has
    source notes; the first section page also has its confidence reason.
    `chip=False` leaves out the section's confidence chip."""
    s = slide(prs); bg_solid(s, WHITE)
    title(s, heading)
    if chip:
        _confidence_chip(s, ctx.get("conf"))
    ctx["page"] += 1
    footer(s, str(ctx["page"]), ctx["project"], ctx["date"])
    notes = ctx.get("notes")
    references = ctx.get("source_notes") or ""
    if ctx.get("section_pages", 0) == 0 and notes:
        note_text = notes
        if references and references not in note_text:
            base_references = ctx.get("base_source_notes") or ""
            if base_references and base_references in note_text:
                note_text = note_text.replace(base_references, references, 1)
            else:
                note_text += "\n\n" + references
    else:
        note_text = references
    if note_text:
        s.notes_slide.notes_text_frame.text = note_text
    ctx["section_pages"] = ctx.get("section_pages", 0) + 1
    return s


# ---------------------------------------------------------------------- covers
# Client requirement (deck 18.09.2026): every RFQ Insight deck states on its
# cover that it contains AI-generated content.
COVER_DISCLAIMER = ("This document includes AI-generated content. "
                    "Please verify important information before use.")


def main_cover(prs, plan, ctx):
    s = slide(prs); bg_image(s, GRAD_COVER())
    pic(s, logo_white_stream(), M, 0.55, 1.7, 1.7 / LOGO_AR)
    lines = [ln.strip() for ln in
             str(plan.get("projectName", "Project Title")).split("\n") if ln.strip()]
    size = 34 if len(lines) <= 1 else 28
    text(s, lines, M, 2.75, PAGE_W - 2 * M, 2.1, size=size, bold=True,
         color=WHITE, font=HF)
    opp = plan.get("oppId", "")
    if opp:
        text(s, opp, M, 4.95, 8, 0.5, size=18, color=WHITE, font=HF)
    text(s, "RFQ Insight", M, 5.65, 6, 0.4, size=16, color=WHITE, font=HF)
    text(s, f"Update {display_date(plan.get('date', ''))}".rstrip(), M, 6.1, 6,
         0.35, size=11, color=WHITE)
    # Between the date label (ends at 6.45") and "AVL · Internal" (starts at
    # 6.95"): one centred line, never overlapping either.
    text(s, COVER_DISCLAIMER, M, PAGE_H - 0.98, PAGE_W - 2 * M, 0.3, size=10,
         color=WHITE, align=PP_ALIGN.CENTER)
    text(s, "AVL  ·  Internal", PAGE_W - M - 4, PAGE_H - 0.55, 4, 0.3,
         size=10, color=WHITE, align=PP_ALIGN.RIGHT)
    ctx["page"] += 1


def toc(prs, plan, ctx):
    s = content_slide(prs, ctx, "Table of Contents")
    qs = plan.get("questions", [])
    half = (len(qs) + 1) // 2 if len(qs) > 8 else len(qs)
    colw = (PAGE_W - 2 * M - 0.6) / (2 if len(qs) > 8 else 1)
    for i, q in enumerate(qs):
        col = 0 if i < half else 1
        row = i if i < half else i - half
        x = M + col * (colw + 0.6)
        y = 1.6 + row * 0.52
        rect(s, x, y, 0.4, 0.4, fill=LIGHT2, rounded=True, radius=0.25)
        # The position, not q["number"]: the optional Additional Information
        # section (number 13) is listed before Source Material (number 12).
        text(s, str(i + 1), x, y + 0.02, 0.4, 0.36, size=13,
             bold=True, color=BLUE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, q.get("titleEn", ""), x + 0.55, y, colw - 0.6, 0.42, size=13,
             color=INK, anchor=MSO_ANCHOR.MIDDLE)


def end_cover(prs, plan, ctx):
    s = slide(prs); bg_image(s, GRAD_COVER())
    w = 2.4
    pic(s, logo_white_stream(), PAGE_W / 2 - w / 2, 2.5, w, w / LOGO_AR)
    text(s, "End of RFQ Insight", 0, 3.9, PAGE_W, 0.7, size=24, bold=True,
         color=WHITE, font=HF, align=PP_ALIGN.CENTER)
    text(s, plan.get("projectName", ""), 0, 4.6, PAGE_W, 0.4, size=13,
         color=WHITE, align=PP_ALIGN.CENTER)
    text(s, "AVL  ·  Internal", 0, PAGE_H - 0.6, PAGE_W, 0.35, size=10,
         color=WHITE, align=PP_ALIGN.CENTER)


# ------------------------------------------------------------------- callout
def callout(s, co, y=None):
    if not (co and co.get("text")):
        return
    if _wrap_count(co["text"], PAGE_W - 2 * M - 1.85, 11) > 3:
        s._rfq_callout_overflow = dict(co)
        co = {"tag": co.get("tag", "NOTE"),
              "text": "Full details follow on continuation slides."}
    y = y if y is not None else CONTENT_BOTTOM - 0.7
    rect(s, M, y, PAGE_W - 2 * M, 0.7, fill=LIGHT, line=LINE, line_w=1)
    tag = co.get("tag", "NOTE")
    tag_fill = {"ASSESSMENT": AMBER, "ANALYSIS": AMBER, "GAP": WARN}.get(tag, BLUE)
    rect(s, M + 0.15, y + 0.17, 1.3, 0.36, fill=tag_fill, rounded=True, radius=0.5)
    text(s, tag, M + 0.15, y + 0.17, 1.3, 0.36, size=8.5, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, co.get("text", ""), M + 1.62, y, PAGE_W - 2 * M - 1.85, 0.7,
         size=11, italic=True, color=INK_SOFT, anchor=MSO_ANCHOR.MIDDLE)


def _callout_continuations(prs, co, heading, ctx):
    """Paginate exact overflow wording instead of spilling out of the diagram band."""
    width, size, max_lines = PAGE_W - 2 * M, 13, 20
    chunks, remaining = [], str(co["text"])
    while remaining:
        low, high = 1, len(remaining)
        while low < high:
            middle = (low + high + 1) // 2
            if _wrap_count(remaining[:middle], width, size) <= max_lines:
                low = middle
            else:
                high = middle - 1
        end = low
        if end < len(remaining):
            boundary = max(remaining.rfind(" ", 0, end), remaining.rfind("\n", 0, end)) + 1
            if boundary > end // 2:
                end = boundary
        chunks.append(remaining[:end])
        remaining = remaining[end:]
    for index, chunk in enumerate(chunks, 1):
        s = content_slide(prs, ctx, heading + " — " + str(co.get("tag", "NOTE"))
                          + ("" if index == 1 else " (cont.)"))
        text(s, chunk, M, 1.55, width, CONTENT_BOTTOM - 1.55,
             size=size, color=INK, line_spacing=1.1)


# ------------------------------------------------------------- content layouts
def L_bullets(prs, d, ctx):
    """Bullet list at FIXED type. Overflow adds a "(cont.)" slide — the old
    auto-shrink ladder had no pagination at all and silently ran the last
    bullets off the bottom of the slide."""
    items = [b.get("text", "") + ("  (inferred)" if b.get("inferred") else "")
             for b in d.get("bullets", [])]
    tw = PAGE_W - 2 * M
    avail = CONTENT_BOTTOM - 1.55                # no note band on a data slide
    pages = _paginate(items,
                      lambda t: _bullet_h(t, tw - BULLET_MARK_W,
                                          BULLET_PT, BULLET_GAP), avail)
    for i, chunk in enumerate(pages):
        s = content_slide(prs, ctx,
                          d.get("title", "") + ("" if i == 0 else " (cont.)"))
        bullets(s, chunk, M, 1.55, tw, avail, size=BULLET_PT, gap=BULLET_GAP)


def _std_table(s, cols, rows, x, y, w, cell_size=10.5, owner_map=None,
               col_widths=None, owner_cols=None, min_row=0.26):
    ncols, nrows = len(cols), len(rows) + 1
    t = s.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w),
                           Inches(min(0.42 * nrows, CONTENT_BOTTOM - y))).table
    t.first_row = False; t.horz_banding = False
    # Declared row heights act as MINIMUMS when rendered: PowerPoint and
    # LibreOffice grow a row to fit its text but never shrink it below the
    # declared value. The default equal split (~0.42"/row) made short pages
    # (e.g. "(cont.)" slides with few rows) render far taller than the
    # pagination estimate and run into the callout band — so declare each
    # row SMALL and let the text alone dictate the real height. `min_row`
    # must stay in sync with the pagination estimate in _row_height.
    t.rows[0].height = Inches(HEADER_MIN)
    for ri in range(1, nrows):
        t.rows[ri].height = Inches(min_row)
    for ci in range(ncols):
        t.columns[ci].width = Inches(col_widths[ci] if col_widths else w / ncols)
    for ci, h in enumerate(cols):
        cell = t.cell(0, ci); cell.fill.solid(); cell.fill.fore_color.rgb = C(BLUE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.05)
        p = cell.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(h)
        r.font.bold = True; r.font.size = Pt(cell_size + 0.5)
        r.font.name = HF; r.font.color.rgb = C(WHITE)
    for ri, row in enumerate(rows, start=1):
        for ci in range(ncols):
            cell = t.cell(ri, ci)
            val = str(row[ci]) if ci < len(row) else ""
            fill, fg, bold = (WHITE if ri % 2 else LIGHT2), INK, False
            if owner_map and ci > 0 and (owner_cols is None or ci in owner_cols):
                oc = owner_map(val)
                if oc:
                    fill, fg, bold = oc, WHITE, True
            cell.fill.solid(); cell.fill.fore_color.rgb = C(fill)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            if owner_map and ci > 0 and fg == WHITE \
                    and (owner_cols is None or ci in owner_cols):
                p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = val
            r.font.size = Pt(cell_size); r.font.name = BF
            r.font.bold = bold; r.font.color.rgb = C(fg)
    return t


# --------------------------------------------------- table height estimation
# Per-character advance widths in EMs (multiply by point size / 72 to get
# inches). Generated from real sans metrics and widened by VERDANA_FACTOR —
# Verdana runs slightly wider than the generic sans used for measurement.
# This replaces the old flat chars-per-line heuristic, which assumed every
# character was ~0.76 em wide and therefore split tables onto a new slide
# when the current slide was barely half full (client feedback: tables must
# fill the whole slide with rows before a continuation slide is added).
_CHAR_EMS = {
    "a": .613, "b": .635, "c": .550, "d": .635, "e": .615, "f": .352,
    "g": .635, "h": .634, "i": .278, "j": .278, "k": .579, "l": .278,
    "m": .974, "n": .634, "o": .612, "p": .635, "q": .635, "r": .411,
    "s": .521, "t": .392, "u": .634, "v": .592, "w": .818, "x": .592,
    "y": .592, "z": .525,
    "A": .684, "B": .686, "C": .698, "D": .770, "E": .632, "F": .575,
    "G": .775, "H": .752, "I": .295, "J": .295, "K": .656, "L": .557,
    "M": .863, "N": .748, "O": .787, "P": .603, "Q": .787, "R": .695,
    "S": .635, "T": .611, "U": .732, "V": .684, "W": .989, "X": .685,
    "Y": .611, "Z": .685,
    " ": .318, ".": .318, ",": .318, ":": .337, ";": .337, "!": .401,
    "?": .531, "'": .275, '"': .460, "(": .390, ")": .390, "[": .390,
    "]": .390, "/": .337, "\\": .337, "|": .337, "-": .361, "–": .500,
    "—": 1.00, "_": .500, "+": .838, "=": .838, "<": .838, ">": .838,
    "%": .950, "&": .780, "#": .838, "@": 1.00, "*": .500, "·": .318,
    "≤": .838, "≥": .838, "±": .838, "°": .500, "€": .636, "…": 1.00,
}
_DIGIT_EM = .636          # all digits are tabular-width
_DEFAULT_EM = .650        # unknown latin-ish glyph
_CJK_EM = 1.05            # full-width CJK glyph
VERDANA_FACTOR = 1.05     # Verdana vs measurement font, incl. safety
BOLD_FACTOR = 1.11        # bold advance widths vs regular
LINE_H_EM = 1.22          # single-spacing line height in EMs
CELL_V_PAD = 0.055        # cell top+bottom margins incl. rendering slack
MIN_ROW_H = 0.26          # matches the declared (minimum) data-row height
# `_std_table` declares its header row at this height (Inches(HEADER_MIN));
# every height ESTIMATE (`_table_h`, `_paginate_rows`) must use the same
# value, or a table that the estimate says fits can still run into the
# footer once python-pptx renders the real, taller header.
HEADER_MIN = 0.36

# ---------------------------------------------------------- fixed type sizes
# Client requirement: type is FIXED per layout family — a table on slide 6 and
# a table on slide 16 render at the same size, always. Content never rescales
# to fit; when it exceeds the slide the layout PAGINATES instead (_paginate).
# That is the whole reason the old auto-shrink ladders are gone: they made two
# slides out of a table that overflowed its budget by 1.7%.
TABLE_PT = 9.5                    # table + ownership cells (header at +0.5)
PANEL_PT, PANEL_GAP = 11, 7       # risk-panel bullets
BULLET_PT, BULLET_GAP = 13, 9     # bullets-layout bullets
BULLET_MARK_W = 0.30              # width the "▪  " marker takes off the column


def _char_em(ch):
    if ch.isdigit():
        return _DIGIT_EM
    if ord(ch) >= 0x2E80:                     # CJK and full-width forms
        return _CJK_EM
    return _CHAR_EMS.get(ch, _DEFAULT_EM)


def _text_w(s, size, bold=False):
    """Rendered width (inches) of one line of text at `size` pt."""
    em = size / 72.0 * VERDANA_FACTOR * (BOLD_FACTOR if bold else 1.0)
    return sum(_char_em(c) for c in s) * em


def _wrap_count(text, avail_w, size, bold=False):
    """Number of lines `text` wraps to inside `avail_w` inches — a greedy
    word-wrap simulation of what PowerPoint does, not a character count."""
    lines = 0
    for seg in str(text).split("\n"):
        words = seg.split()
        if not words:
            lines += 1
            continue
        cur = 0.0
        lines += 1
        space_w = _text_w(" ", size, bold)
        for word in words:
            ww = _text_w(word, size, bold)
            if ww > avail_w:                      # long token: hard-break
                if cur > 0:
                    lines += 1
                lines += max(int(ww // avail_w), 0)
                cur = ww % avail_w
                continue
            need = ww if cur == 0 else cur + space_w + ww
            if need > avail_w and cur > 0:
                lines += 1
                cur = ww
            else:
                cur = need
    return max(lines, 1)


def _row_height(row, col_widths, cell_size, min_row=MIN_ROW_H, bold=False):
    """Estimated rendered height of one table row. PowerPoint grows a row to
    fit its wrapped text but never shrinks it below the declared height, so
    the estimate is max(declared minimum, text height)."""
    lines = 1
    for ci, cell in enumerate(row):
        w = col_widths[ci] if ci < len(col_widths) else col_widths[-1]
        avail = max(w - 0.15, 0.3)                # cell L/R margins
        lines = max(lines, _wrap_count(cell, avail, cell_size, bold))
    return max(min_row, lines * LINE_H_EM * cell_size / 72.0 + CELL_V_PAD)


def _paginate(items, h_of, budget, overhead=0.0):
    """Greedy split of `items` into pages whose MEASURED heights fit `budget`.

    `overhead` is the per-page chrome every page pays for (a table's repeated
    header row). Shared by the table, bullets and panel layouts, so all three
    split on what the text actually measures — never on an item count."""
    chunks, cur, used = [], [], overhead
    for it in items:
        h = h_of(it)
        if cur and used + h > budget:
            chunks.append(cur); cur, used = [], overhead
        cur.append(it); used += h
    if cur or not chunks:
        chunks.append(cur)

    # anti-thin: even out the final page (a 1-2 item orphan looks broken) by
    # pulling items back from the previous page — but only while the final
    # page still fits its budget, so the footer band stays clear.
    while (len(chunks) > 1 and len(chunks[-1]) + 2 <= len(chunks[-2])
           and overhead + sum(map(h_of, chunks[-1]))
               + h_of(chunks[-2][-1]) <= budget):
        chunks[-1].insert(0, chunks[-2].pop())
    return chunks


def _paginate_rows(rows, col_widths, cell_size, budget, header=None,
                   min_row=MIN_ROW_H):
    """Split table rows into pages that fit the vertical budget.

    `header` (the column-title row, if given) is measured like a data row —
    long per-phase headers wrap to 2-3 lines, so a fixed allowance
    under-reserves and the last row can run off the slide."""
    hdr_h = (_row_height(header, col_widths, cell_size + 0.5,
                         min_row=HEADER_MIN, bold=True) if header else 0.42)
    return _paginate(rows,
                     lambda r: _row_height(r, col_widths, cell_size,
                                           min_row=min_row),
                     budget, overhead=hdr_h)


def _bullet_h(t, avail_w, size, gap):
    """Measured height of one bullet's line block at `size` pt, plus its gap."""
    return _wrap_count(t, avail_w, size) * size * LINE_H_EM / 72.0 + gap / 72.0


def L_table(prs, d, ctx):
    rows = d.get("rows", [])
    cols = d.get("columns", [])
    tw = PAGE_W - 2 * M
    widths = [tw / max(len(cols), 1)] * max(len(cols), 1)
    # The whole slide below the title belongs to the table: a data slide never
    # reserves a note band (client: notes and gaps belong in Risk Detection).
    bottom = CONTENT_BOTTOM
    chunks = _paginate_rows(rows, widths, TABLE_PT, bottom - 1.55, header=cols)
    for i, ch in enumerate(chunks):
        s = content_slide(prs, ctx,
                          d.get("title", "") + ("" if i == 0 else " (cont.)"))
        _std_table(s, cols, ch, M, 1.55, tw, cell_size=TABLE_PT)


OPP_WIDTH_FRACTION = 0.62    # left-aligned tables, ~60% wide (client slide)
OPP_GAP = 0.12               # vertical gap between the three tables (keeps the normal 10+2+4 rows on one slide)


def _table_h(cols, rows, widths, size, header_min=HEADER_MIN):
    """Estimated rendered height of a whole table, header included."""
    head = _row_height(cols, widths, size + 0.5, min_row=header_min, bold=True)
    return head + sum(_row_height(r, widths, size) for r in rows)


def L_opportunity(prs, d, ctx):
    """Opportunity Info: the main table, then Products (%) and Skills (%)
    stacked under it (client format 18.09.2026). A table that would run
    into the footer moves to an "(cont.)" slide instead of shrinking.

    The main table's first column holds short labels ("Customer", "KAM
    (PTE)", ...); a Products/Skills table's first column holds a full
    product or skill name instead, so each gets its own label/value split
    rather than sharing one - three headed tables (plus the two gaps
    between them) already spend more of the vertical budget than the
    single-table layouts, so the main table's column split and the top
    margin are tuned to leave the client's Opportunity Info content on one
    slide whenever it fits (the normal 10+2+4 rows do); a table that does
    not paginates onto "(cont.)" slides instead of overrunning the footer
    (never truncates a row)."""
    from .opportunity import PRODUCTS_COLUMNS, SKILLS_COLUMNS
    tw = (PAGE_W - 2 * M) * OPP_WIDTH_FRACTION
    main_widths = [tw * 0.32, tw * 0.68]
    pct_widths = [tw * 0.40, tw * 0.60]
    top, bottom = 1.3, CONTENT_BOTTOM
    blocks = [(d.get("columns") or [], d.get("rows") or [], main_widths),
              (PRODUCTS_COLUMNS, d.get("products") or [], pct_widths),
              (SKILLS_COLUMNS, d.get("skills") or [], pct_widths)]
    blocks = [(cols, rows, w) for cols, rows, w in blocks if cols and rows]
    title_base = d.get("title", "")
    state = {"s": None, "y": top, "page": 0}

    def _new_slide():
        state["s"] = content_slide(prs, ctx, title_base
                                   + ("" if state["page"] == 0 else " (cont.)"))
        state["page"] += 1
        state["y"] = top

    for cols, rows, widths in blocks:
        h = _table_h(cols, rows, widths, TABLE_PT)
        if state["s"] is None or (state["y"] + h > bottom and state["y"] > top):
            _new_slide()
        if state["y"] + h > bottom:
            # Does not fit even alone at the top of an empty slide: paginate
            # its rows onto "(cont.)" slides instead of truncating or
            # overrunning the footer - every row/cell still appears.
            chunks = _paginate_rows(rows, widths, TABLE_PT,
                                    bottom - state["y"], header=cols)
            for i, chunk in enumerate(chunks):
                if i > 0:
                    _new_slide()
                chunk_h = _table_h(cols, chunk, widths, TABLE_PT)
                _std_table(state["s"], cols, chunk, M, state["y"], tw,
                          cell_size=TABLE_PT, col_widths=widths)
                state["y"] += chunk_h + OPP_GAP
        else:
            _std_table(state["s"], cols, rows, M, state["y"], tw,
                      cell_size=TABLE_PT, col_widths=widths)
            state["y"] += h + OPP_GAP
    if state["s"] is None:             # nothing to show: keep the section
        content_slide(prs, ctx, title_base)


def _panel(s, x, heading, items, accent):
    w = (PAGE_W - 2 * M - 0.4) / 2
    rect(s, x, 1.55, w, 4.55, fill=LIGHT, line=LINE, line_w=1)
    rect(s, x, 1.55, w, 0.62, fill=accent)
    text(s, heading, x + 0.25, 1.55, w - 0.5, 0.62, size=14, bold=True,
         color=WHITE, font=HF, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, items, x + 0.25, 2.45, w - 0.5, 3.5, size=12, gap=9,
            marker_color=accent)
    return x + w + 0.4


def L_twoColumn(prs, d, ctx):
    s = content_slide(prs, ctx, d.get("title", ""))
    nx = _panel(s, M, d.get("leftHeading", ""), d.get("leftItems", []), BLUE)
    _panel(s, nx, d.get("rightHeading", ""), d.get("rightItems", []), CYAN)


# Accent per well-known panel heading, so e.g. Project Risks and Document
# Risks keep distinct colors even on separate slides.
_PANEL_ACCENTS = (("document", CYAN), ("project", BLUE))


def L_panel(prs, d, ctx):
    """ONE full-width headed panel per slide — the client-confirmed Risk
    Detection format (Project Risks and Document Risks on separate slides).

    Type is FIXED at PANEL_PT. A group that genuinely outgrows one slide gets
    a "(cont.)" slide; it is never pre-split by item count, which is what put
    8 risks onto two half-empty slides."""
    heading = d.get("heading", "")
    items = d.get("items", [])
    # A follow-up question's heading is the user's own wording: one accent.
    accent = BLUE if ctx.get("number") == ADDITIONAL_INFO_ID else next(
        (c for k, c in _PANEL_ACCENTS if k in heading.lower()), BLUE)
    x, w = M, PAGE_W - 2 * M
    top = 1.55
    bottom = CONTENT_BOTTOM                  # no note band on a risk slide
    avail = bottom - top - 0.62 - 0.55
    tw = w - 0.5 - BULLET_MARK_W
    pages = _paginate(items,
                      lambda t: _bullet_h(t, tw, PANEL_PT, PANEL_GAP), avail)
    for i, chunk in enumerate(pages):
        suffix = "" if i == 0 else " (cont.)"
        s = content_slide(prs, ctx, d.get("title", "") + suffix)
        rect(s, x, top, w, bottom - top, fill=LIGHT, line=LINE, line_w=1)
        rect(s, x, top, w, 0.62, fill=accent)
        text(s, heading + suffix, x + 0.25, top, w - 0.5, 0.62, size=14,
             bold=True, color=WHITE, font=HF, anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, chunk, x + 0.25, top + 0.9, w - 0.5, avail, size=PANEL_PT,
                gap=PANEL_GAP, marker_color=accent)


def L_prosCons(prs, d, ctx):
    s = content_slide(prs, ctx, d.get("title", ""))
    nx = _panel(s, M, "Advantages", d.get("pros", []), GREEN)
    _panel(s, nx, "Considerations", d.get("cons", []), WARN)


def L_kpi(prs, d, ctx):
    s = content_slide(prs, ctx, d.get("title", ""))
    kpis = d.get("kpis", [])[:4]
    n = max(len(kpis), 1)
    kw = (PAGE_W - 2 * M - (n - 1) * 0.3) / n
    ky, kh = 2.0, 3.2
    accents = [BLUE, CYAN, GREEN, LIME]
    for i, k in enumerate(kpis):
        x = M + i * (kw + 0.3)
        rect(s, x, ky, kw, kh, fill=LIGHT, line=LINE, line_w=1)
        rect(s, x, ky, kw, 0.08, fill=accents[i % 4])
        val = str(k.get("value", "")) + ((" " + k["unit"]) if k.get("unit") else "")
        vsize = 30 if len(val) <= 8 else (22 if len(val) <= 12 else 17)
        text(s, val, x + 0.1, ky + 0.8, kw - 0.2, 1.0, size=vsize, bold=True,
             color=BLUE, font=HF, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             wrap=False)
        text(s, k.get("label", ""), x + 0.15, ky + 2.0, kw - 0.3, 1.0, size=11,
             color=INK_SOFT, align=PP_ALIGN.CENTER)


def L_prose(prs, d, ctx):
    s = content_slide(prs, ctx, d.get("title", ""))
    if d.get("caption"):
        text(s, d["caption"], M, 1.45, PAGE_W - 2 * M, 0.4, size=12, italic=True,
             color=INK_SOFT)
    body = d.get("body", "")
    paras = [p.strip() for p in body.split("\n") if p.strip()] or [body]
    # Auto-shrink so a full 5-paragraph client summary fits ONE slide.
    words = len(body.split())
    size, gap, spacing = ((14, 10, 1.25) if words <= 140 else
                          (13, 8, 1.2) if words <= 200 else
                          (12, 7, 1.15) if words <= 270 else
                          (11, 6, 1.1))
    text(s, paras, M, 1.95, PAGE_W - 2 * M, CONTENT_BOTTOM - 1.95, size=size,
         color=INK, line_spacing=spacing, space_after=gap)


# ------------------------------------------------------------- DIAGRAM layouts
_TL_MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun",
     "jul", "aug", "sep", "oct", "nov", "dec"])}


def _date_key(s):
    """'~Oct 2026' / 'Q1 2028' / 'Sep 2027 / Q1 2028' -> sortable int, or None.

    A month+year pair wins over a quarter, and each is paired with ITS OWN
    year — so dual-format dates like "~Sep 2027 / Q1 2028" key on Sep 2027."""
    import re
    s = str(s).lower()
    m = re.search(r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)"
                  r"[a-z]*\.?,?[\s\-./]*(20\d\d)", s)
    if m:
        return int(m.group(2)) * 100 + _TL_MONTHS[m.group(1)]
    q = re.search(r"q([1-4])\s*[/\- ]?\s*(20\d\d)", s)
    if q:
        return int(q.group(2)) * 100 + int(q.group(1)) * 3 - 1
    ym = re.search(r"(20\d\d)", s)
    if not ym:
        return None
    year = int(ym.group(1))
    q2 = re.search(r"q([1-4])", s)
    if q2:
        return year * 100 + int(q2.group(1)) * 3 - 1
    return year * 100 + 6


def _chrono_sort(steps):
    """Stable chronological sort. Steps without a parseable date inherit the
    previous dated step's position (so '+10 business days' stays glued to its
    anchor milestone); leading undated steps stay first."""
    keyed, last = [], -1
    for i, st in enumerate(steps):
        k = _date_key(st.get("date", ""))
        if k is None:
            k = last
        else:
            last = k
        keyed.append((k, i, st))
    return [st for _, _, st in sorted(keyed, key=lambda t: (t[0], t[1]))]


def L_timeline(prs, d, ctx):
    """Milestone graphic: nodes on a gradient axis, alternating cards.
    Steps are always drawn in chronological order, whatever order the plan
    supplied them in."""
    steps = _chrono_sort(d.get("steps", []))
    # Balanced pagination: max 7 nodes per page, but pages share the steps
    # evenly (8 steps -> 4+4, not 7+1 with an orphaned last milestone).
    n_pages = max(1, -(-len(steps) // 7))
    per = -(-len(steps) // n_pages) if steps else 1
    pages = [steps[i:i + per] for i in range(0, len(steps), per)] or [[]]
    seq = 0
    for pi, chunk in enumerate(pages):
        s = content_slide(prs, ctx,
                          d.get("title", "") + ("" if pi == 0 else " (cont.)"))
        n = max(len(chunk), 1)
        tx0, tx1, baseY = M + 0.8, PAGE_W - M - 0.8, 3.9
        gap = (tx1 - tx0) / max(n - 1, 1)
        s.shapes.add_picture(GRAD_BAR(), Inches(tx0 - 0.3), Inches(baseY - 0.03),
                             Inches(tx1 - tx0 + 0.6), Inches(0.06))
        # Small NUMBERED dots on the axis — the milestone NAME lives in the
        # card as a bold heading, so long labels can never overflow a circle.
        node = 0.34
        card_w, card_h = 1.92, 1.5
        for i, st in enumerate(chunk):
            seq += 1
            cx = tx0 + (i * gap if n > 1 else (tx1 - tx0) / 2)
            shape(s, MSO_SHAPE.OVAL, cx - node / 2, baseY - node / 2, node, node,
                  fill=BLUE, line=WHITE, line_w=2.0)
            text(s, str(seq), cx - node / 2, baseY - node / 2, node, node,
                 size=10, bold=True, color=WHITE, font=HF, wrap=False,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            up = (i % 2 == 0)
            cardY = (baseY - 0.35 - card_h) if up else (baseY + 0.35)
            connector(s, cx, baseY + (-node / 2 if up else node / 2), cx,
                      cardY + (card_h if up else 0), color=GRAY_SOFT, weight=1)
            rect(s, cx - card_w / 2, cardY, card_w, card_h, fill=LIGHT,
                 line=LINE, line_w=1)
            label = str(st.get("label", ""))
            lsize = 9.5 if len(label) <= 34 else 8.5
            text(s, label, cx - card_w / 2 + 0.07, cardY + 0.07,
                 card_w - 0.14, 0.46, size=lsize, bold=True, color=INK,
                 font=HF, align=PP_ALIGN.CENTER)
            text(s, st.get("date", ""), cx - card_w / 2 + 0.07, cardY + 0.55,
                 card_w - 0.14, 0.24, size=9, bold=True, color=BLUE, font=HF,
                 align=PP_ALIGN.CENTER)
            text(s, st.get("desc", ""), cx - card_w / 2 + 0.07, cardY + 0.8,
                 card_w - 0.14, card_h - 0.87, size=8, color=INK_SOFT,
                 align=PP_ALIGN.CENTER)
        if pi == len(pages) - 1:
            callout(s, d.get("callout"))


# ------------------------------------------------------------------- gantt
# Date parsing lives in `dates.py`, NOT here: `rfq/validation.py`'s C5 check has to
# order milestones exactly as this module places them, and it cannot import
# this file without dragging python-pptx into a checker that must stay
# runnable on a report JSON alone. One parser, imported twice.
from .dates import time_pos as _time_pos  # noqa: E402


# ------------------------------------------------- gantt layout constants (in)
# Inline milestone labels: candidate box widths, tried narrowest-first so a
# short label claims only the room it needs and its neighbours keep theirs.
LAB_W_STEPS = (0.95, 1.10, 1.25, 1.40, 1.55, 1.75)
LAB_PAD = 0.10            # min horizontal gap between two boxes in one lane
LANE_GAP = 0.08           # vertical gap between lanes
LANE_CLR_UP = 0.34        # baseline -> bottom of the nearest ABOVE lane
LANE_CLR_DN = 0.30        # baseline -> top of the nearest BELOW lane
MAX_LANES = 4
# (date pt, label pt, max label lines) — biggest readable rung first
LAB_LADDER = ((9.0, 8.5, 2), (8.5, 8.0, 2), (8.0, 7.5, 2), (7.5, 7.0, 3))

# Legend mode: numbered markers on the axis + a keyed list below the chart.
BADGE_D = 0.24            # numbered marker diameter
BADGE_PITCH = 0.27        # vertical pitch of stacked badge rows
BADGE_PAD = 0.03          # min horizontal gap between two badges in a row
BADGE_ROWS = 3
BADGE_LEAD = 0.10         # baseline -> first badge row
LEG_COL_GAP = 0.28
LEG_NUM_W = 0.26
LEG_MAX_COLS = 6
LEG_LADDER = (10.0, 9.0, 8.0, 7.0)

# Client-confirmed: the numbered-badge + keyed legend IS the timing format, so
# the chart reads the same at 3 milestones and at 24. Set False to restore the
# adaptive inline-label path (_plan_milestone_labels below).
GANTT_ALWAYS_LEGEND = True

# Phase-bar vertical budget.
BAR_GAP_MAX, BAR_GAP_MIN = 0.10, 0.04
BAR_H_MAX, BAR_H_MIN = 0.42, 0.14
BAR_NAME_MIN_H = 0.20     # below this a bar's name goes outside the bar
BASE_CLR = 0.30           # last bar -> milestone baseline
MS_RESERVE_MIN = 1.20     # milestone zone is never thinner than this


def _ms_box(st, dsize, lsize, max_lines):
    """(width, height) of ONE milestone's inline label block, measured.

    The date is one bold line; the name may wrap to `max_lines`. Width is the
    narrowest candidate at which the name wraps within that budget, floored at
    the date line so the date never overflows its own box."""
    lbl = str(st.get("label", ""))
    w = LAB_W_STEPS[-1]
    for cand in LAB_W_STEPS:
        if _wrap_count(lbl, cand, lsize, bold=True) <= max_lines:
            w = cand
            break
    w = max(w, _text_w(str(st.get("date", "")), dsize, bold=True))
    h = (dsize * LINE_H_EM / 72.0) + 0.03 + max_lines * (lsize * LINE_H_EM / 72.0)
    return w, h


def _lane_slots(line_y, lab_h, body_top, body_bot, bars_bot=None):
    """Label lanes as (side, y_top), derived from the REAL free space rather
    than a fixed count — a shorter label box buys more lanes. Lanes above the
    baseline only exist in the room the phase bars leave free. Nearest to the
    baseline first, so labels stay close to their markers."""
    up_top = (bars_bot + 0.10) if bars_bot is not None else body_top
    pitch = lab_h + LANE_GAP
    n_up = int(max(line_y - LANE_CLR_UP - up_top + LANE_GAP, 0) // pitch)
    n_dn = int(max(body_bot - line_y - LANE_CLR_DN + LANE_GAP, 0) // pitch)
    lanes = [("up", line_y - LANE_CLR_UP - lab_h - k * pitch) for k in range(n_up)]
    lanes += [("down", line_y + LANE_CLR_DN + k * pitch) for k in range(n_dn)]
    lanes.sort(key=lambda l: abs(l[1] - line_y))
    return lanes[:MAX_LANES]


def _pack_lanes(items, n_lanes, x_lo, x_hi, pad=LAB_PAD):
    """Greedy first-fit packing of `items` = [(centre_x, width)] ascending by
    centre, into `n_lanes` horizontal lanes. Returns [(lane_idx, left_x)], or
    None if ANY item cannot be placed without overlapping its lane neighbour.

    The test is on the CLAMPED interval, not the raw centre distance: a label
    at the axis edge gets shifted inward, and two centres far enough apart can
    still collide once both are clamped. First-fit is exact here because the
    centres ascend and the clamp is monotone, so the last box placed in a lane
    is always that lane's binding constraint.

    Returning None is the point: the caller escalates (smaller type, then the
    numbered-legend layout) instead of stacking labels on top of each other."""
    right = [-1e6] * max(n_lanes, 1)
    out = []
    for cx, w in items:
        lx = min(max(cx - w / 2.0, x_lo), x_hi - w)
        li = next((i for i in range(n_lanes) if lx >= right[i] + pad), None)
        if li is None:
            return None
        right[li] = lx + w
        out.append((li, lx))
    return out


def _spread_badges(centres, n_rows, x_lo, x_hi, pad=BADGE_PAD, d=BADGE_D):
    """Place numbered markers on `n_rows` rows below the baseline, as close to
    their true time position as possible. Returns [(row, left_x)].

    Each marker takes the row that lets it sit nearest its true position, and
    is nudged right only as far as its row neighbour forces. Markers therefore
    never overlap, and a leader line back to the axis (drawn by the caller)
    keeps the true date unambiguous when one has been nudged."""
    right = [x_lo - pad] * max(n_rows, 1)
    out = []
    for cx in centres:
        want = min(max(cx - d / 2.0, x_lo), x_hi - d)
        r = min(range(len(right)), key=lambda i: (max(want, right[i] + pad), i))
        lx = max(want, right[r] + pad)
        right[r] = lx + d
        out.append((r, lx))
    return out


def _plan_milestone_labels(placed, X, line_y, body_top, body_bot,
                           x_lo, x_hi, bars_bot=None):
    """Try to label every milestone directly on the chart.

    OFF BY DEFAULT: GANTT_ALWAYS_LEGEND short-circuits this, so the chart uses
    legend mode whatever the milestone count. Kept for the flag's False path.

    Returns (dsize, lsize, lanes, [(lane_idx, left_x, w, lab_h)]) or None,
    which means "labels cannot be placed cleanly — use legend mode".

    Each font rung re-derives the lane count, because smaller type means
    shorter boxes means more lanes; the biggest readable type that fits wins."""
    for dsize, lsize, max_lines in LAB_LADDER:
        boxes = [_ms_box(st, dsize, lsize, max_lines) for _, st in placed]
        lab_h = max(h for _, h in boxes)
        lanes = _lane_slots(line_y, lab_h, body_top, body_bot, bars_bot)
        if not lanes:
            continue
        got = _pack_lanes([(X(t), w) for (t, _), (w, _) in zip(placed, boxes)],
                          len(lanes), x_lo, x_hi)
        if got is not None:
            return dsize, lsize, lanes, [(li, lx, w, lab_h)
                                         for (li, lx), (w, _) in zip(got, boxes)]
    return None


def _legend_geom(entries, f, usable_w):
    """(row_h, date_w, entry_w, ncols) for the legend at `f` pt. Columns are
    maximised: more columns means fewer rows means a shorter legend, which
    hands the vertical space back to the chart itself."""
    row_h = max(f * LINE_H_EM / 72.0 + 0.06, 0.20)
    date_w = max(_text_w(dt, f, bold=True) for _, dt, _ in entries) + 0.06
    name_w = max(_text_w(nm, f) for _, _, nm in entries)
    entry_w = LEG_NUM_W + 0.08 + date_w + 0.10 + name_w
    ncols = max((c for c in range(1, min(LEG_MAX_COLS, len(entries)) + 1)
                 if c * entry_w + (c - 1) * LEG_COL_GAP <= usable_w), default=1)
    return row_h, date_w, entry_w, ncols


def _legend_fit(entries, usable_w, max_h):
    """Size the milestone legend at the largest readable font that fits.
    `entries` = [(num, date, name)]. Returns
    (fsize, ncols, nrows, row_h, date_w, entry_w), or None if even the
    smallest rung overflows `max_h`."""
    if not entries:
        return None
    for f in LEG_LADDER:
        row_h, date_w, entry_w, ncols = _legend_geom(entries, f, usable_w)
        rows = -(-len(entries) // ncols)
        if rows * row_h + 0.06 <= max_h:
            return f, ncols, rows, row_h, date_w, entry_w
    return None


def _legend_capacity(entries, usable_w, max_h):
    """How many entries the legend can hold at its smallest rung — the basis
    for the spill note when a plan carries an implausible milestone count."""
    row_h, _, _, ncols = _legend_geom(entries, LEG_LADDER[-1], usable_w)
    return max(ncols * int((max_h - 0.06) // row_h), 1)


def _fit_bars(n, avail):
    """Bar height + gap that fit `n` phase bars into `avail` inches."""
    n = max(n, 1)
    gap = BAR_GAP_MAX if n <= 8 else BAR_GAP_MIN
    bar_h = min(BAR_H_MAX, (avail - (n - 1) * gap) / n)
    if bar_h < BAR_H_MIN:
        gap = BAR_GAP_MIN
        bar_h = max((avail - (n - 1) * gap) / n, 0.09)
    return bar_h, gap


def _gantt_baseline(n, body_top, body_bot):
    """Provisional bar block + baseline. Bars take at most 55% of the body and
    always leave MS_RESERVE_MIN for the milestones — the old formula had a
    0.24" floor per bar, so past ~10 bars the block ran off the slide."""
    body_h = body_bot - body_top
    avail = min(body_h * 0.55, body_h - MS_RESERVE_MIN - BASE_CLR)
    bar_h, gap = _fit_bars(n, avail)
    bars_h = n * bar_h + (n - 1) * gap
    return bar_h, gap, body_top + bars_h, body_top + bars_h + BASE_CLR


def _gantt_plan(placed, bars_in, X, body_top, body_bot, x_lo, x_hi):
    """Decide the whole body layout WITHOUT drawing: bar height, baseline y,
    and the numbered-badge + legend layout (or, only when GANTT_ALWAYS_LEGEND
    is off and they fit, inline milestone labels).

    Returns a dict; `spill` is the milestones that did not fit even the
    smallest legend rung (the caller discloses them in the callout)."""
    n_bars = len(bars_in)
    if bars_in:
        bar_h, bar_gap, bars_bot, line_y = _gantt_baseline(n_bars, body_top, body_bot)
    else:
        bar_h, bar_gap, bars_bot = 0.42, 0.0, None
        line_y = (body_top + body_bot) / 2

    inline = None if GANTT_ALWAYS_LEGEND else _plan_milestone_labels(
        placed, X, line_y, body_top, body_bot, x_lo, x_hi, bars_bot)
    if inline is not None:
        return {"bar_h": bar_h, "bar_gap": bar_gap, "line_y": line_y,
                "inline": inline, "badges": None, "legend": None, "spill": []}

    # ---- legend mode. Badges keep their true time position but de-collide
    # over up to BADGE_ROWS rows; the legend then claims the height it needs
    # and the baseline (and with it the bar block) is pushed up to make room.
    badges = _spread_badges([X(t) for t, _ in placed], BADGE_ROWS, x_lo, x_hi)
    badge_band = BADGE_LEAD + (max(r for r, _ in badges) + 1) * BADGE_PITCH

    entries = [(str(i), str(st.get("date", "")), str(st.get("label", "")))
               for i, (_, st) in enumerate(placed, start=1)]
    min_bars_h = (n_bars * BAR_H_MIN + (n_bars - 1) * BAR_GAP_MIN) if bars_in else 0.42
    max_leg_h = body_bot - (body_top + min_bars_h + BASE_CLR + badge_band)
    usable_w = PAGE_W - 2 * M

    spill = []
    fit = _legend_fit(entries, usable_w, max_leg_h)
    if fit is None:                    # pathological count: keep what fits
        cap = _legend_capacity(entries, usable_w, max_leg_h)
        spill = entries[cap:]
        entries = entries[:cap]
        fit = _legend_fit(entries, usable_w, max_leg_h)
    if fit is None:                    # body has no room at all for a legend
        return {"bar_h": bar_h, "bar_gap": bar_gap, "line_y": line_y,
                "inline": None, "badges": badges, "badge_band": badge_band,
                "legend": None, "entries": [], "spill": entries}

    fsize, ncols, nrows, row_h, date_w, entry_w = fit
    leg_h = nrows * row_h + 0.06
    line_y = body_bot - leg_h - badge_band
    if bars_in:
        bar_h, bar_gap = _fit_bars(n_bars, line_y - BASE_CLR - body_top)
    return {"bar_h": bar_h, "bar_gap": bar_gap, "line_y": line_y,
            "inline": None, "badges": badges, "badge_band": badge_band,
            "legend": (fsize, ncols, nrows, row_h, date_w, entry_w),
            "entries": entries, "spill": spill}


def _zone_caption(s, label, y0, y1, color=INK_SOFT):
    """Rotated swimlane caption in the left gutter, spanning y0..y1 — labels
    the PHASES zone and the MILESTONES zone of the Gantt so the two sections
    read as clearly separated bands (client feedback)."""
    h = max(y1 - y0, 0.6)
    cx, cy = 0.32, (y0 + y1) / 2.0
    tb = text(s, label, cx - h / 2, cy - 0.14, h, 0.28, size=8, bold=True,
              color=color, font=HF, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    tb.rotation = 270
    return tb


def _draw_legend(s, plan, line_y):
    """Numbered milestone legend below the chart, filled column-major so the
    reading order stays chronological."""
    fsize, ncols, nrows, row_h, date_w, entry_w = plan["legend"]
    total_w = ncols * entry_w + (ncols - 1) * LEG_COL_GAP
    x0 = M + (PAGE_W - 2 * M - total_w) / 2
    top = line_y + plan["badge_band"]
    bd = min(row_h - 0.02, LEG_NUM_W)
    for i, (num, date_txt, name) in enumerate(plan["entries"]):
        ex = x0 + (i // nrows) * (entry_w + LEG_COL_GAP)
        ey = top + (i % nrows) * row_h
        shape(s, MSO_SHAPE.OVAL, ex, ey + (row_h - bd) / 2, bd, bd, fill=BLUE)
        text(s, num, ex, ey + (row_h - bd) / 2, bd, bd,
             size=max(fsize - 2.0, 6.0), bold=True, color=WHITE, font=HF,
             wrap=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        dx = ex + LEG_NUM_W + 0.08
        text(s, date_txt, dx, ey, date_w, row_h, size=fsize, bold=True,
             color=BLUE, font=HF, wrap=False, anchor=MSO_ANCHOR.MIDDLE)
        nx = dx + date_w + 0.10
        text(s, name, nx, ey, max(ex + entry_w - nx, 0.3), row_h, size=fsize,
             color=INK, wrap=False, anchor=MSO_ANCHOR.MIDDLE)


def L_gantt(prs, d, ctx):
    """One-slide Gantt chart: year/quarter header grid, the project (or
    per-phase) bar(s) spanning their real calendar range, and every dated
    milestone as a numbered badge at its proportional position, keyed to a
    legend below the chart — the timing picture the sales team asked for,
    everything in ONE diagram and the same shape at any milestone count
    (GANTT_ALWAYS_LEGEND). Undated milestones are disclosed in the callout;
    if nothing is datable at all, falls back to the classic milestone
    timeline."""
    steps = d.get("steps", [])
    placed = [(t, st) for st in steps
              if (t := _time_pos(st.get("date", ""))) is not None]
    if len(placed) < 2:
        return L_timeline(prs, d, ctx)
    placed = [(pos, idx, st) for idx, (pos, st) in enumerate(placed)]
    placed.sort(key=lambda p: (round(p[0] * 15) / 15, p[1]))
    placed = [(pos, st) for pos, _, st in placed]
    unplaced = [st for st in steps if _time_pos(st.get("date", "")) is None]

    bars_in = []
    for b in (d.get("bars") or []):
        t0, t1 = _time_pos(b.get("start", "")), _time_pos(b.get("end", ""), end=True)
        if t0 is not None and t1 is not None and t1 > t0:
            bars_in.append((t0, t1, b.get("name", "")))
    # Chronological bar order regardless of plan order — the drawn diagram is
    # a pure function of the dates, never of model output ordering.
    bars_in.sort(key=lambda b: (b[0], b[1], b[2]))

    # merge undated milestones into the callout so nothing silently vanishes
    co = dict(d.get("callout") or {})
    if unplaced:
        note = "Not on time axis (no calendar date): " + "; ".join(
            f"{st.get('label', '?')} ({st.get('date', '')})" for st in unplaced)
        co = {"tag": co.get("tag", "NOTE"),
              "text": ((co.get("text", "").rstrip(". ") + ". ")
                       if co.get("text") else "") + note}
    has_callout = bool(co.get("text"))

    s = content_slide(prs, ctx, d.get("title", ""))

    # ---- time axis: pad to whole quarters
    lo = min([p[0] for p in placed] + [b[0] for b in bars_in])
    hi = max([p[0] for p in placed] + [b[1] for b in bars_in])
    q0 = int(lo // 3) * 3
    q1 = int(-(-(hi + 0.35) // 3)) * 3            # exclusive end, padded
    n_q = max((q1 - q0) // 3, 1)

    ax0, ax1 = M + 0.05, PAGE_W - M - 0.05
    ax_w = ax1 - ax0

    def X(t):
        return ax0 + (t - q0) / (q1 - q0) * ax_w

    y0, yr_h, q_h = 1.45, 0.3, 0.26
    grid_top = y0 + yr_h + q_h
    x_lo, x_hi = M - 0.1, PAGE_W - M + 0.1

    # ---- measure and decide the whole body BEFORE drawing any of it. A
    # legend that spills adds a callout, which shrinks the body — so re-plan
    # once against the smaller body rather than drawing over the callout.
    for _ in range(2):
        bottom = (CONTENT_BOTTOM - 0.7 - 0.12) if has_callout \
            else CONTENT_BOTTOM
        body_top, body_bot = grid_top + 0.1, bottom - 0.1
        plan = _gantt_plan(placed, bars_in, X, body_top, body_bot, x_lo, x_hi)
        if not plan["spill"]:
            break
        note = (f"{len(plan['spill'])} further milestones "
                f"({plan['spill'][0][2]} … {plan['spill'][-1][2]}) are listed "
                f"on the Timing table slide.")
        co = {"tag": co.get("tag", "NOTE"),
              "text": ((co.get("text", "").rstrip(". ") + ". ")
                       if co.get("text") else "") + note}
        if has_callout:
            break
        has_callout = True
    bar_h, bar_gap, line_y = plan["bar_h"], plan["bar_gap"], plan["line_y"]

    # ---- header: year cells spanning their quarters, quarter cells below
    years = sorted({(q0 + 3 * i) // 12 for i in range(n_q)})
    for yr in years:
        a = max(yr * 12, q0); b = min((yr + 1) * 12, q1)
        rect(s, X(a), y0, X(b) - X(a), yr_h, fill=BLUE)
        text(s, str(yr), X(a), y0, X(b) - X(a), yr_h, size=11, bold=True,
             color=WHITE, font=HF, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    q_label = n_q <= 16                            # label quarters if they fit
    for i in range(n_q):
        a = q0 + 3 * i
        rect(s, X(a), y0 + yr_h, X(a + 3) - X(a), q_h, fill=LIGHT2, line=LINE,
             line_w=0.75)
        if q_label:
            text(s, f"Q{(a % 12) // 3 + 1}", X(a), y0 + yr_h,
                 X(a + 3) - X(a), q_h, size=8.5, color=INK_SOFT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # quarter gridlines through the chart body (year boundaries darker)
    for i in range(n_q + 1):
        a = q0 + 3 * i
        connector(s, X(a), grid_top, X(a), bottom,
                  color=(GRAY_SOFT if a % 12 == 0 else LINE),
                  weight=(1.2 if a % 12 == 0 else 0.75))

    # ---- bars
    if bars_in:
        accents = [BLUE, CYAN, GREEN, LIME, GRAY, GRAY_SOFT]
        for i, (t0, t1, name) in enumerate(bars_in):
            by = body_top + i * (bar_h + bar_gap)
            bx0, bx1 = X(t0), X(t1)
            shape(s, MSO_SHAPE.PENTAGON, bx0, by, max(bx1 - bx0, 0.3), bar_h,
                  fill=accents[i % len(accents)])
            nsize = next((c for c in (10, 9, 8) if
                          _text_w(name, c, bold=True) <= bx1 - bx0 - 0.35), 8)
            # A name only goes inside the bar when the bar is both wide AND
            # tall enough — a 0.14" bar at 12 phases cannot hold 8pt type.
            if _text_w(name, 8, bold=True) <= bx1 - bx0 - 0.35 \
                    and bar_h >= BAR_NAME_MIN_H:
                text(s, name, bx0 + 0.1, by, bx1 - bx0 - 0.3, bar_h,
                     size=nsize, bold=True, color=WHITE, font=HF,
                     anchor=MSO_ANCHOR.MIDDLE)
            else:                                  # name outside, to the right
                nh = max(bar_h, 0.16)
                text(s, name, min(bx1 + 0.08, ax1 - 1.6),
                     by + (bar_h - nh) / 2, 1.6, nh,
                     size=8, bold=True, color=INK, font=HF,
                     anchor=MSO_ANCHOR.MIDDLE)
        bars_bot_y = body_top + len(bars_in) * bar_h \
            + (len(bars_in) - 1) * bar_gap
        _zone_caption(s, "PHASES", body_top, bars_bot_y)
    else:
        t0, t1 = placed[0][0], placed[-1][0]
        bh = 0.42
        # Inline mode keeps the bar straddling the baseline; in legend mode the
        # baseline has been pushed down to make room, so the bar takes the
        # CENTRE of the project zone above it (hugging the baseline left a
        # large dead band under the header).
        by = (line_y - bh / 2) if plan["inline"] \
            else (body_top + (line_y - 0.12 - body_top) / 2 - bh / 2)
        label = d.get("barLabel") or str(ctx["project"]).split("\n")[0]
        shape(s, MSO_SHAPE.PENTAGON, X(t0), by,
              max(X(t1) - X(t0), 0.5), bh, fill=BLUE)
        lsize = next((c for c in (13, 11, 9) if
                      _text_w(label, c, bold=True) <= X(t1) - X(t0) - 0.4), 9)
        text(s, label, X(t0) + 0.15, by, X(t1) - X(t0) - 0.4, bh,
             size=lsize, bold=True, color=WHITE, font=HF,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if not plan["inline"]:
            _zone_caption(s, "PROJECT", by, by + bh)

    if bars_in or not plan["inline"]:      # milestone baseline
        s.shapes.add_picture(GRAD_BAR(), Inches(ax0), Inches(line_y - 0.02),
                             Inches(ax_w), Inches(0.04))

    if plan["inline"]:
        # ---- direct labels: arrow marker on the baseline, name in a packed
        # lane. The packer guaranteed every box fits, so nothing overlaps.
        dsize, lsize, lanes, boxes = plan["inline"]
        for (t, st), (li, lx, lw, lh) in zip(placed, boxes):
            cx = X(t)
            side, ly_ = lanes[li]
            aw, ah = 0.16, 0.2
            if side == "up":
                shape(s, MSO_SHAPE.DOWN_ARROW, cx - aw / 2,
                      line_y - 0.21 - ah + 0.02, aw, ah, fill=BLUE)
                y_conn0, y_conn1 = ly_ + lh, line_y - 0.21 - ah
            else:
                shape(s, MSO_SHAPE.UP_ARROW, cx - aw / 2, line_y + 0.21 - 0.02,
                      aw, ah, fill=BLUE)
                y_conn0, y_conn1 = line_y + 0.21 + ah, ly_
            if abs(y_conn1 - y_conn0) > 0.12:
                connector(s, cx, y_conn0, cx, y_conn1, color=GRAY_SOFT,
                          weight=0.75)
            dh = dsize * LINE_H_EM / 72.0
            text(s, str(st.get("date", "")), lx, ly_, lw, dh, size=dsize,
                 bold=True, color=BLUE, font=HF, align=PP_ALIGN.CENTER)
            text(s, str(st.get("label", "")), lx, ly_ + dh + 0.03, lw,
                 lh - dh - 0.03, size=lsize, bold=True, color=INK,
                 align=PP_ALIGN.CENTER)
    else:
        # ---- legend mode: the axis carries numbered badges only, so labels
        # can never collide however many milestones there are.
        for i, ((t, _), (row, bx)) in enumerate(zip(placed, plan["badges"]),
                                                start=1):
            cx = X(t)
            by = line_y + BADGE_LEAD + row * BADGE_PITCH
            if row or abs(bx + BADGE_D / 2 - cx) > 0.02:
                connector(s, cx, line_y, bx + BADGE_D / 2, by,
                          color=GRAY_SOFT, weight=0.75)
            shape(s, MSO_SHAPE.OVAL, bx, by, BADGE_D, BADGE_D, fill=BLUE,
                  line=WHITE, line_w=1.25)
            text(s, str(i), bx, by, BADGE_D, BADGE_D,
                 size=8.0 if i < 10 else 7.0, bold=True, color=WHITE, font=HF,
                 wrap=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if plan["legend"]:
            _draw_legend(s, plan, line_y)
        # The milestone zone (baseline + badges + legend) is its own labelled
        # band, cleanly separated from the phase bars above it.
        _zone_caption(s, "MILESTONES", line_y, body_bot)
    callout(s, co if has_callout else None)


PH_MAX_PER_ROW = 6        # chevrons per row before the diagram wraps
PH_ROW_GAP = 0.18
PH_OVERLAP = 0.25


def _ellipsize(s, avail_w, size, max_lines):
    """Trim `s` on word boundaries until it wraps within `max_lines`."""
    words = str(s).split()
    while words and _wrap_count(" ".join(words) + "…", avail_w, size) > max_lines:
        words.pop()
    return (" ".join(words) + "…") if words else ""


def L_phases(prs, d, ctx):
    """Development phases as a chevron process diagram.

    Wraps onto further rows past PH_MAX_PER_ROW rather than dropping phases:
    the deck must show every in-scope phase the answer names, however many
    that is. Only descriptions are ever shortened, and visibly."""
    s = content_slide(prs, ctx, d.get("title", ""))
    phases = d.get("phases", [])
    n = max(len(phases), 1)
    co = dict(d.get("callout") or {})
    accents = [BLUE, CYAN, GREEN, LIME, GRAY, GRAY_SOFT]

    rows = -(-n // PH_MAX_PER_ROW)
    per = -(-n // rows)                       # balanced: 7 -> 4+3, 12 -> 6+6
    w = (PAGE_W - 2 * M + PH_OVERLAP * (per - 1)) / per
    if rows == 1:                             # unchanged single-row geometry
        top, band, h, desc_h = 2.2, None, 1.15, 1.9
    else:
        top = 1.75
        bottom = (CONTENT_BOTTOM - 0.7 - 0.15) if co.get("text") \
            else CONTENT_BOTTOM
        band = (bottom - top) / rows
        h = min(1.15, band * 0.42)
        desc_h = band - h - 0.22 - PH_ROW_GAP

    # Auto-shrink the name font so the longest (client-verbatim) phase name
    # still fits its chevron — phase names are copied from the RFQ's own
    # naming and can be long. Names are never truncated.
    name_w = w - 0.55 - 0.55
    longest = max((str(p.get("name", "")) for p in phases), key=len, default="")
    nsize = 8
    for cand in (15, 13, 11.5, 10, 9, 8):
        nsize = cand
        if _wrap_count(longest, name_w, cand, bold=True) \
                * cand * LINE_H_EM / 72.0 <= h * 0.88:
            break

    def _desc(p):
        t = str(p.get("desc", ""))
        return (t + "  ·  HW milestone").strip("  ·") if p.get("hw") else t

    # On a wrapped diagram the description uses the whole column pitch rather
    # than the chevron's text inset — worth ~0.75" of width per column.
    desc_w = (w - PH_OVERLAP - 0.10) if rows > 1 else (w - 0.55 - 0.4)
    dsize = 8
    for cand in (10.5, 9.5, 8.5, 8):
        dsize = cand
        if max((_wrap_count(_desc(p), desc_w, cand) for p in phases), default=1) \
                * cand * LINE_H_EM / 72.0 <= desc_h:
            break
    max_lines = max(int(desc_h / (dsize * LINE_H_EM / 72.0)), 1)

    clipped = False
    for i, ph in enumerate(phases):
        r, c = divmod(i, per)
        chunk = min(per, n - r * per)
        row_w = chunk * w - PH_OVERLAP * (chunk - 1)
        x = M + (PAGE_W - 2 * M - row_w) / 2 + c * (w - PH_OVERLAP)
        y = top if band is None else top + r * band
        cv = shape(s, MSO_SHAPE.CHEVRON, x, y, w, h, fill=accents[i % len(accents)])
        try:
            cv.adjustments[0] = 0.35
        except Exception:
            pass
        inset = 0.55 if c else 0.35
        text(s, ph.get("name", ""), x + inset, y, w - inset - 0.55, h,
             size=nsize, bold=True, color=WHITE, font=HF,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        desc = _desc(ph)
        if _wrap_count(desc, desc_w, dsize) > max_lines:
            desc = _ellipsize(desc, desc_w, dsize, max_lines)
            clipped = True
        dx = (x + PH_OVERLAP / 2) if rows > 1 else (x + inset)
        text(s, desc, dx, y + h + 0.25, desc_w, desc_h, size=dsize,
             color=INK_SOFT)

    if clipped:
        note = ("Phase descriptions abbreviated to fit the diagram — full text "
                "in the Development Phases answer.")
        co = {"tag": co.get("tag", "NOTE"),
              "text": ((co.get("text", "").rstrip(". ") + ". ")
                       if co.get("text") else "") + note}
    callout(s, co if co.get("text") else None)


def L_collab(prs, d, ctx):
    parties = d.get("parties") or []
    chunks = [parties[i:i + 6] for i in range(0, len(parties), 6)] or [[]]
    for index, chunk in enumerate(chunks):
        page = dict(d)
        page["parties"] = chunk
        page["title"] = d.get("title", "") + ("" if index == 0 else " (cont.)")
        _draw_collab_page(prs, page, ctx)


def _draw_collab_page(prs, d, ctx):
    """Hub-and-spoke collaboration model: project hub + self-sizing party nodes.

    Node heights expand to fit their role text (estimated line count at the
    node font size), so long roles never overflow the box.
    """
    s = content_slide(prs, ctx, d.get("title", ""))
    hub = d.get("hub", ctx["project"])
    parties = d.get("parties", [])
    import math
    n = max(len(parties), 1)
    node_w = 2.9
    role_size = 8.0
    cpl = max(int((node_w - 0.24) / (0.0105 * role_size)), 10)  # chars per line

    def node_height(pt):
        role = str(pt.get("role", ""))
        lines = -(-len(role) // cpl) if role else 0
        return min(0.34 + 0.12 + lines * (role_size + 4) / 72.0 + 0.08, 2.1)

    heights = [node_height(pt) for pt in parties]
    has_callout = bool(d.get("callout") and d["callout"].get("text"))
    top_area = 1.35
    # when a callout exists, the whole ring stays ABOVE the callout band
    bottom_limit = (CONTENT_BOTTOM - 0.7 - 0.12) if has_callout \
        else CONTENT_BOTTOM
    cxx = PAGE_W / 2
    cyy = (top_area + bottom_limit) / 2 + 0.05
    # vertical radius adapts so the tallest top/bottom nodes stay inside bounds
    max_h = max(heights) if heights else 1.0
    ry = (bottom_limit - top_area) / 2 - max_h / 2 - 0.05
    rx = (PAGE_W - 2 * M - node_w) / 2 - 0.1
    hw, hh = 2.6, 0.95
    positions = []
    for i in range(n):
        ang = -math.pi / 2 + i * (2 * math.pi / n)
        positions.append((cxx + rx * math.cos(ang), cyy + ry * math.sin(ang)))
    type_color = {"customer": BLUE, "avl": GREEN, "supplier": GRAY,
                  "third": CYAN}
    for (px, py), pt in zip(positions, parties):
        connector(s, cxx, cyy, px, py, color=GRAY_SOFT, weight=1.5,
                  dashed=pt.get("dashed", False))
    for (px, py), pt, nh in zip(positions, parties, heights):
        col = type_color.get(str(pt.get("type", "")).lower(), CYAN)
        x0, y0 = px - node_w / 2, py - nh / 2
        rect(s, x0, y0, node_w, nh, fill=WHITE, line=col, line_w=2.0,
             rounded=True, radius=0.12)
        rect(s, x0, y0, node_w, 0.34, fill=col, rounded=False)
        nm = pt.get("name", "")
        nsize = 10
        for cand in (10, 9, 8, 7):
            nsize = cand
            if len(nm) * 0.0105 * cand <= node_w - 0.16:
                break
        text(s, nm, x0 + 0.08, y0, node_w - 0.16, 0.34,
             size=nsize, bold=True, color=WHITE, font=HF,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, pt.get("role", ""), x0 + 0.12, y0 + 0.42, node_w - 0.24,
             nh - 0.5, size=role_size, color=INK)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cxx - hw / 2, cyy - hh / 2, hw, hh,
          fill=BLUE)
    text(s, hub, cxx - hw / 2 + 0.1, cyy - hh / 2, hw - 0.2, hh, size=11.5,
         bold=True, color=WHITE, font=HF, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    callout(s, d.get("callout"))


def L_ownership(prs, d, ctx):
    """Scope-of-Work matrix with color-coded owner cells + legend."""
    abbr = str(d.get("customerAbbr", "Customer")).lower()

    def owner_map(v):
        vl = v.strip().lower()
        if vl == "avl":
            return OWNER_COLORS["avl"]
        if vl in ("customer", abbr):
            return OWNER_COLORS["customer"]
        if vl == "supplier":
            return OWNER_COLORS["supplier"]
        if vl == "joint":
            return OWNER_COLORS["joint"]
        return None

    cols = d.get("columns", [])
    rows = d.get("rows", [])
    tw = PAGE_W - 2 * M
    widths = [tw / max(len(cols), 1)] * max(len(cols), 1)
    if len(cols) > 1:
        widths = [tw * 0.46] + [tw * 0.54 / (len(cols) - 1)] * (len(cols) - 1)
    ly = CONTENT_BOTTOM - 0.30                       # legend line
    bottom = ly - 0.15                               # no note band, just legend
    # owner columns = non-stage columns (a "Phases"/"Stage" column is info,
    # not an owner); denser packing so ~18 key-task rows fit one slide like
    # the client reference deck
    stage_hdrs = {"phase", "phases", "stage", "development stage",
                  "development phase", "dev stage", "development stages"}
    owner_cols = {i for i in range(1, len(cols))
                  if str(cols[i]).strip().lower() not in stage_hdrs}
    chunks = _paginate_rows(rows, widths, TABLE_PT, bottom - 1.5,
                            header=cols, min_row=0.21)
    for i, ch in enumerate(chunks):
        s = content_slide(prs, ctx,
                          d.get("title", "") + ("" if i == 0 else " (cont.)"))
        _std_table(s, cols, ch, M, 1.5, tw, cell_size=TABLE_PT,
                   owner_map=owner_map, col_widths=widths,
                   owner_cols=owner_cols, min_row=0.21)
        # legend
        lx = M
        for label, key in (("AVL", "avl"),
                           (d.get("customerAbbr", "Customer"), "customer"),
                           ("Supplier", "supplier"),
                           ("Joint — to be split", "joint")):
            rect(s, lx, ly, 0.28, 0.28, fill=OWNER_COLORS[key])
            text(s, label, lx + 0.36, ly - 0.02, 1.6, 0.32, size=10,
                 color=INK_SOFT, anchor=MSO_ANCHOR.MIDDLE)
            lx += 0.5 + max(0.9, 0.12 * len(label))


# -------------------------------------------------------- product pictures
CAPTION_H = 0.36          # caption box under each picture (2 lines at 10 pt)
SOURCE_TEXT_H = 0.24      # source box under the caption (1 line at 8 pt)
PIC_GAP = 0.4             # gap between picture cells


def _fit(aspect, slot_w, slot_h):
    """Largest (w, h) with w / h == aspect that fits slot_w x slot_h."""
    if aspect <= 0 or slot_w <= 0 or slot_h <= 0:
        return max(slot_w, 0), max(slot_h, 0)
    if slot_w / slot_h > aspect:
        return slot_h * aspect, slot_h
    return slot_w, slot_w / aspect


def _picture_slots(n, top, bottom):
    """(x, y, w, h) picture boxes for 1-4 pictures; caption and source lines
    sit below each box, all above `bottom`. 1 = one large box, 2 = side by
    side, 3-4 = a 2 x 2 grid (3 leaves the last cell empty)."""
    full_w = PAGE_W - 2 * M
    label = CAPTION_H + SOURCE_TEXT_H
    if n <= 1:
        return [(M, top, full_w, bottom - top - label)]
    half_w = (full_w - PIC_GAP) / 2
    if n == 2:
        return [(M + i * (half_w + PIC_GAP), top, half_w, bottom - top - label)
                for i in range(2)]
    cell_h = (bottom - top - PIC_GAP) / 2 - label
    cells = [(M + c * (half_w + PIC_GAP), top + r * (cell_h + label + PIC_GAP),
              half_w, cell_h) for r in range(2) for c in range(2)]
    return cells[:n]


def L_pictures(prs, d, ctx):
    """Product pictures copied from the RFQ, aspect ratio kept, each with its
    caption and source. A picture that cannot be loaded leaves its caption
    and source in place with an empty box; the run report lists it. No
    confidence chip: the app adds these illustrations, no answer is rated.
    The picture disclaimer belongs to the section's first-page notes."""
    s = content_slide(prs, ctx, d.get("title", ""), chip=False)
    items = d.get("pictures", [])[:4]
    top = 1.55
    for (x, y, w, h), item in zip(_picture_slots(len(items), top, CONTENT_BOTTOM),
                                  items):
        path = resolve_picture_file(item, ctx.get("images_dir"))
        loaded = deck_picture(path)
        if loaded is not None:
            data, iw, ih = loaded
            pw, ph = _fit(iw / ih if ih else 0, w, h)
            pic(s, io.BytesIO(data), x + (w - pw) / 2, y + (h - ph) / 2, pw, ph)
        text(s, item.get("caption", ""), x, y + h + 0.04, w, CAPTION_H - 0.04,
             size=10, color=INK, align=PP_ALIGN.CENTER)
        text(s, item.get("source", ""), x, y + h + CAPTION_H, w, SOURCE_TEXT_H,
             size=8, italic=True, color=INK_SOFT, align=PP_ALIGN.CENTER)


LAYOUTS = {
    "bullets": L_bullets, "table": L_table, "twoColumn": L_twoColumn,
    "panel": L_panel, "opportunity": L_opportunity,
    "prosCons": L_prosCons, "timeline": L_timeline, "kpi": L_kpi,
    "prose": L_prose,
    # diagram layouts
    "phases": L_phases, "collab": L_collab, "ownership": L_ownership,
    "gantt": L_gantt, "pictures": L_pictures,
}


# ------------------------------------------------------------------------ build
def create_presentation(plan, images_dir=None):
    plan = copy.deepcopy(plan)
    prs = Presentation()
    prs.slide_width = Inches(PAGE_W); prs.slide_height = Inches(PAGE_H)
    ctx = {"project": plan.get("projectName", "Project Title"),
           "date": plan.get("date", ""), "page": 0}
    # Kept for raw layout tools; client decks exclude picture appendices.
    ctx["images_dir"] = images_dir
    main_cover(prs, plan, ctx)
    toc(prs, plan, ctx)
    for q in plan.get("questions", []):
        ctx["conf"] = q.get("confidence") or ""
        ctx["notes"] = (q.get("confidenceNotes") or "").partition(
            "\n\nHow the rating works:")[0].replace("Why this rating:", "Reason:")
        if not ctx["conf"] and not ctx["notes"]:
            ctx["notes"] = q.get("sourceDetails") or ""
        ctx["section_pages"] = 0
        ctx["number"] = q.get("number")
        for sd in q.get("slides", []):
            if sd.get("layout") == "pictures":
                continue
            ctx["source_notes"] = (sd.get("sourceDetails") or
                                   (sd.get("sourceLine") if q.get("number") == ADDITIONAL_INFO_ID else "") or
                                   q.get("sourceDetails") or q.get("sourceLine") or "")
            ctx["base_source_notes"] = ctx["source_notes"]
            # Saved plans sometimes put provenance in an inventory column.
            # Move that column intact into notes before displaying the table.
            if q.get("number") == SOURCES and sd.get("layout") == "table":
                columns = sd.get("columns", [])
                moved = [i for i, col in enumerate(columns) if col.strip().casefold()
                         in ("corpus locations", "source locations", "pages", "page references")]
                if moved:
                    entries = ["File: " + row[0] + "\n  Location: " + "; ".join(
                        row[i] for i in moved if i < len(row)) for row in sd.get("rows", []) if row]
                    ctx["source_notes"] += "\n\nSource inventory:\n" + "\n\n".join(entries)
                    sd["columns"] = [col for i, col in enumerate(columns) if i not in moved]
                    sd["rows"] = [[value for i, value in enumerate(row) if i not in moved]
                                  for row in sd.get("rows", [])]
            # Only explicit page/slide annotations are removed. Standard codes,
            # technical clause identifiers, targets and document names survive.
            annotations = []
            def without_page_annotations(value):
                if isinstance(value, str):
                    # Explicit abbreviated page citations only. Full words
                    # such as 'Page 1 must be signed' may be deliverables.
                    number = r"\d+(?!\d|\.\d)(?:\s*[-\u2010-\u2015]\s*\d+(?!\d|\.\d))?"
                    pattern = r"(?P<context>\b(?:RFP|RFQ|SOR|SOW|offer|agreement|document|same offer|its)\s+)(?:p\.|pp\.)\s+" + number + r"(?:\s*,\s*" + number + r")*"
                    def remember(match):
                        annotations.append(value)
                        return match.group("context")
                    clean = re.sub(pattern, remember, value, flags=re.IGNORECASE)
                    # A file extension immediately precedes a source locator.
                    file_pattern = r"(?P<context>\.(?:pdf|docx|pptx|xlsx)\s+)(?:p\.|pp\.|slides?)\s+" + number + r"(?:\s*,\s*" + number + r")*"
                    clean = re.sub(file_pattern, remember, clean, flags=re.IGNORECASE)
                    clean = re.sub(r" {2,}", " ", clean).strip()
                    if clean != value:
                        clean = re.sub(r"\s+([.,;:)])", r"\1", clean)
                    return clean
                if isinstance(value, list):
                    return [without_page_annotations(item) for item in value]
                if isinstance(value, dict):
                    return {key: (item if key in ("layout", "sourceLine", "sourceDetails")
                                  else without_page_annotations(item)) for key, item in value.items()}
                return value
            sd = without_page_annotations(sd)
            if annotations:
                ctx["source_notes"] += "\n\nOriginal source references:\n" + "\n".join(
                    "- " + item for item in dict.fromkeys(annotations))
            if q.get("number") == ADDITIONAL_INFO_ID:
                ctx["notes"] = ctx["source_notes"]
            fn = LAYOUTS.get(sd.get("layout"))
            if fn:
                first = len(prs.slides)
                fn(prs, sd, ctx)
                overflow = [getattr(s, "_rfq_callout_overflow", None)
                            for s in list(prs.slides)[first:]]
                seen = set()
                for co in overflow:
                    if co and co["text"] not in seen:
                        _callout_continuations(prs, co, sd.get("title", ""), ctx)
                        seen.add(co["text"])
            else:
                L_bullets(prs, {"title": sd.get("title", ""),
                                "bullets": [{"text": "Unsupported layout: "
                                             + str(sd.get("layout"))}]}, ctx)
    end_cover(prs, plan, ctx)
    return prs


def render_to_bytes(plan, images_dir=None):
    buf = io.BytesIO()
    create_presentation(plan, images_dir=images_dir).save(buf)
    return buf.getvalue()
