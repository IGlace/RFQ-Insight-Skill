"""Native per-format text extraction.

Replaces the Power Automate OCR path: every file type is parsed with the
library that understands it, so text arrives structured instead of
flattened. The single biggest quality lever here is TABLES - they come out
as real markdown pipe tables, which is the shape the analysis prompt is
told to produce. The prompt's garbled-table reconstruction protocol stays
in place as a safety net.

Images cannot be read here. They are unpacked to bytes, given a stable id
and an inline marker at the position they occurred, and handed upward for a
vision pass (see figures.py). Putting the
transcription back at its original position keeps the surrounding context
and keeps an answer's `sources` naming the parent file.

An image id is `<key>__imgNN`, where `key` identifies the CORPUS ENTRY the
picture came from - not its filename. Everything downstream (the file on
disk, the marker in the text, the transcription filename, the merge) keys
on that id, so it has to be unique across the whole corpus. See image_key.
"""

import io
import os
import re
from dataclasses import dataclass, field

SUPPORTED_EXTENSIONS = {".docx", ".pptx", ".xlsx", ".pdf", ".csv", ".txt",
                        ".png", ".jpg", ".jpeg"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}

# Carried over from the Power Automate corpus contract so the per-file
# metadata keys in corpus.json stay recognisable.
CATEGORY = {
    ".pdf": "pdf", ".png": "image", ".jpg": "image", ".jpeg": "image",
    ".docx": "office_doc", ".pptx": "office_slide", ".xlsx": "office_sheet",
    ".csv": "text_csv", ".txt": "text_plain",
}
METHOD = {
    ".pdf": "pdfplumber", ".docx": "python-docx", ".pptx": "python-pptx",
    ".xlsx": "openpyxl", ".csv": "raw_text", ".txt": "raw_text",
    ".png": "vision", ".jpg": "vision", ".jpeg": "vision",
}

_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")


class UnreadableFile(Exception):
    """A file we support but could not read.

    `reason` is shown to the user and is always plain language. `detail`
    carries the technical cause for the run report and is never surfaced to
    a salesperson. The spec requires both audiences be served separately;
    one string cannot do it.
    """

    def __init__(self, reason: str, detail: str = ""):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


def image_key(rel_path: str) -> str:
    """A readable, filesystem-safe identity for one corpus entry.

    Image ids have to be unique across the CORPUS, not across a folder.
    Two files can share a basename in different subfolders
    (annexes/Spec.docx and attachments/Spec.docx), and an id built from the
    basename alone would give both documents the same id: one image's bytes
    would overwrite the other's on disk, and the merge would hand one
    document a verbatim transcription of a chart it does not contain.

    The path relative to the RFQ folder is unique by construction. Slugging
    it keeps the id usable as a filename and still readable in a message to
    the user - a file at the top of the folder keeps its plain name.
    """
    import re
    slug = re.sub(r"[^A-Za-z0-9_.-]+", "_", rel_path or "")
    return slug.strip("_") or "file"


@dataclass
class RawImage:
    """One image pulled out of a document, before filtering."""
    source: str              # parent filename, e.g. "SOR.docx"
    ordinal: int             # 1-based within the parent
    data: bytes
    ext: str                 # ".png", ".jpeg", ...
    locator: str = ""        # "paragraph 47" / "slide 3" / "page 12" / "standalone"
    key: str = ""            # corpus-unique identity; defaults to `source`
    presized: bool = False   # producer already applied a size judgement
    # sha256 of `data`, set by figures.filter_images so write_images does
    # not hash the same bytes again.
    digest: str = field(default="", compare=False, repr=False)

    def __post_init__(self):
        # `source` stays the DISPLAY name - it is what the user is shown and
        # what an answer's `sources` name. `key` is the identity. They differ
        # only when two corpus entries share a basename.
        if not self.key:
            self.key = self.source

    @property
    def image_id(self) -> str:
        return f"{self.key}__img{self.ordinal:02d}"


@dataclass
class Extracted:
    text: str
    images: list = field(default_factory=list)
    locations: dict = field(default_factory=dict)
    raw_text: str = ""
    extraction_metrics: dict = field(default_factory=dict)


class _ImageCollector:
    def __init__(self, source: str, key: str = ""):
        self.source = source
        self.key = key or source
        self.images: list = []

    def add(self, data: bytes, ext: str, locator: str,
            presized: bool = False) -> RawImage:
        img = RawImage(self.source, len(self.images) + 1, data,
                       (ext or ".png").lower(), locator, self.key, presized)
        self.images.append(img)
        return img


def marker_for(image_id: str) -> str:
    """The placeholder left where an image was, replaced by its
    transcription in figures.py."""
    return f"[[EMBEDDED IMAGE {image_id} — transcription pending]]"


def _why(exc: Exception) -> tuple:
    """(reason for the user, detail for the run report)."""
    detail = f"{type(exc).__name__}: {exc}"
    text = detail.lower()
    if "password" in text or "encrypt" in text:
        return "password-protected", detail
    if "permission" in text or "being used by another" in text:
        return "locked by another program", detail
    if ("not a zip" in text or "badzipfile" in text
            or "file is not a" in text or "no such" in text
            or "cannot be opened" in text or "package not found" in text):
        # "package not found" also fires when the file has vanished between
        # the folder scan and the open. Same advice either way: the file the
        # name points at is not a usable document.
        return "corrupt, or not really the file type its name says", detail
    return "I couldn't open it", detail


def _docx_row_cells(row) -> list:
    """One value per grid position, with a merged cell's text appearing only
    in the first position it spans.

    python-docx returns the same underlying <tc> for every position a merged
    cell covers, so reading .text per position would repeat the value across
    the span and widen the table with phantom data.
    """
    seen, cells = set(), []
    for cell in row.cells:
        key = id(cell._tc)
        if key in seen:
            cells.append("")
        else:
            seen.add(key)
            cells.append(cell.text)
    return cells


def _pptx_row_cells(row) -> list:
    """Same rule as the docx helper: a spanned cell contributes once."""
    return ["" if getattr(cell, "is_spanned", False) else cell.text
            for cell in row.cells]


def rows_to_markdown(rows) -> str:
    """Render table rows as a markdown pipe table. Row 0 is the header.

    Empty rows are dropped, ragged rows padded, embedded newlines collapsed
    and literal pipes escaped, so one cell can never break the table shape.
    """
    import re
    cleaned = []
    for row in rows:
        cells = [re.sub(r"\s+", " ", str("" if c is None else c)).strip()
                 .replace("|", "\\|") for c in row]
        if any(cells):
            cleaned.append(cells)
    if not cleaned:
        return ""
    width = max(len(r) for r in cleaned)
    cleaned = [r + [""] * (width - len(r)) for r in cleaned]
    header, body = cleaned[0], cleaned[1:]
    out = ["| " + " | ".join(header) + " |",
           "| " + " | ".join(["---"] * width) + " |"]
    out.extend("| " + " | ".join(r) + " |" for r in body)
    return "\n".join(out)


# ------------------------------------------------------------------- docx
# Word paragraph styles that mark a heading. Matched on the style NAME, which
# is localised: an English template says "Heading 2", a German one
# "Überschrift 2". "Title"/"Titel" is the document title style.
_HEADING_PREFIXES = ("heading", "überschrift")
_TITLE_STYLES = {"title", "titel"}


def _paragraph_properties(paragraph):
    """Direct properties first, then the complete basedOn style chain."""
    if paragraph._p.pPr is not None:
        yield paragraph._p.pPr
    try:
        style = paragraph.style
        seen = set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            if style.element.pPr is not None:
                yield style.element.pPr
            style = style.base_style
    except (AttributeError, KeyError):
        return


def _heading_level(paragraph):
    """Word's outline is authoritative, including custom/inherited styles.

    An explicit body level (9) overrides an inherited heading level. TOC
    entries are navigation, not new occurrences of the source section.
    """
    from docx.oxml.ns import qn
    try:
        name = (paragraph.style.name or "").strip().casefold()
    except (AttributeError, KeyError):
        name = ""
    if name.startswith(("toc", "table of contents")):
        return None
    for props in _paragraph_properties(paragraph):
        outline = props.find(qn("w:outlineLvl"))
        if outline is not None:
            try:
                value = int(outline.get(qn("w:val")))
                return value + 1 if 0 <= value < 9 else None
            except (TypeError, ValueError):
                return None
    if name in _TITLE_STYLES:
        return 0
    if name.startswith(_HEADING_PREFIXES):
        match = re.search(r"(\d+)\s*$", name)
        return int(match.group(1)) if match else 1
    return None


def _is_heading(paragraph) -> bool:
    return _heading_level(paragraph) is not None


def _docx_blocks(parent):
    """Visit paragraphs/tables once, including content controls, in order."""
    for child in parent.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag in ("p", "tbl"):
            yield child
        elif tag in ("sdt", "sdtContent", "customXml", "ins"):
            yield from _docx_blocks(child)


class _WordNumbering:
    """Reproduce supported decimal heading labels from Word numbering XML.

    Never invent numbers for bullets, unsupported number formats, or absent
    parent counters. Those headings retain their actual text instead.
    """
    def __init__(self, document):
        try:
            self.root = document.part.numbering_part.element
        except (AttributeError, KeyError):
            self.root = None
        self.counts = {}

    def label(self, paragraph):
        from docx.oxml.ns import qn
        values = {}
        for props in _paragraph_properties(paragraph):
            numpr = props.find(qn("w:numPr"))
            if numpr is not None:
                for name in ("numId", "ilvl"):
                    node = numpr.find(qn("w:" + name))
                    if node is not None and name not in values:
                        values[name] = node.get(qn("w:val"))
        if self.root is None or values.get("numId", "0") == "0":
            return ""
        try:
            numid, level = values["numId"], int(values.get("ilvl", "0"))
            num = next(n for n in self.root.findall(qn("w:num"))
                       if n.get(qn("w:numId")) == numid)
            abstractid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
            abstract = next(n for n in self.root.findall(qn("w:abstractNum"))
                            if n.get(qn("w:abstractNumId")) == abstractid)
            levels = {int(n.get(qn("w:ilvl"))): n
                      for n in abstract.findall(qn("w:lvl"))}
            starts = {}
            for override in num.findall(qn("w:lvlOverride")):
                index = int(override.get(qn("w:ilvl")))
                definition = override.find(qn("w:lvl"))
                if definition is not None:
                    levels[index] = definition
                start = override.find(qn("w:startOverride"))
                if start is not None:
                    starts[index] = int(start.get(qn("w:val")))
            definition = levels[level]
            if definition.find(qn("w:numFmt")).get(qn("w:val")) != "decimal":
                return ""
            pattern = definition.find(qn("w:lvlText")).get(qn("w:val"))
            counters = self.counts.setdefault(numid, {})
            startnode = definition.find(qn("w:start"))
            start = starts.get(level, int(startnode.get(qn("w:val")))
                               if startnode is not None else 1)
            counters[level] = counters.get(level, start - 1) + 1
            for index in list(counters):
                if index > level:
                    del counters[index]
            referenced = [int(n) - 1 for n in re.findall(r"%(\d+)", pattern)]
            if any(index not in counters for index in referenced):
                return ""
            return re.sub(r"%(\d+)", lambda m: str(counters[int(m.group(1)) - 1]),
                          pattern)
        except (StopIteration, AttributeError, KeyError, TypeError, ValueError):
            return ""


_MANUAL_SECTION = re.compile(r"^(\d{1,2}(?:\.\d{1,2})*)(?![\d.])(?:\s*[^\d\W]|\s*[\u3400-\u9fff])", re.UNICODE)


def _word_page_blocks(path, blocks, pagination):
    """Accept only a native map for these exact bytes and ordered blocks."""
    import hashlib
    try:
        if (not isinstance(pagination, dict) or
                pagination.get("method") != "Microsoft Word rendered pagination" or
                pagination.get("source_sha256") != _file_bytes_hash(path)):
            return None
        pages, mapped = pagination.get("pages"), pagination.get("blocks")
        if type(pages) is not int or pages < 1 or not isinstance(mapped, list) or len(mapped) != len(blocks):
            return None
        previous = 1
        for child, item in zip(blocks, mapped):
            kind = "paragraph" if child.tag.rsplit("}", 1)[-1] == "p" else "table"
            if not isinstance(item, dict) or item.get("kind") != kind:
                return None
            first, last = item.get("start_page"), item.get("end_page")
            if type(first) is not int or type(last) is not int or not previous <= first <= last <= pages:
                return None
            previous = first
        return mapped
    except (OSError, TypeError, ValueError):
        return None


def _file_bytes_hash(path):
    import hashlib
    with open(path, "rb") as source:
        return hashlib.sha256(source.read()).hexdigest()


def render_word_pages(path, work_dir, timeout=60):
    """Best-effort native pagination of a private copy, with bounded runtime.

    The helper lives beside the plugin assets and uses Word only on Windows.
    Portable callers retain chapter references when native Word is unavailable.
    """
    import hashlib
    import json
    from pathlib import Path
    import shutil
    import subprocess
    import tempfile
    if os.name != "nt":
        return None
    shell = shutil.which("pwsh") or shutil.which("powershell")
    helper = Path(__file__).resolve().parents[2] / "assets" / "word_pages.ps1"
    if not shell or not helper.is_file():
        return None
    try:
        task = Path(tempfile.mkdtemp(prefix="word-pages-", dir=work_dir))
        copied, output = task / "source.docx", task / "pages.json"
        shutil.copyfile(path, copied)
        proc = subprocess.Popen([shell, "-NoProfile", "-NonInteractive", "-File",
                                 str(helper), "-Source", str(copied), "-Output", str(output)],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        try:
            if proc.wait(timeout=timeout) != 0 or not output.is_file():
                return None
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait()
            # The helper records only its own hidden Word process, never an
            # existing user's Office process. Retain scratch evidence.
            pid_file = task / "word.pid"
            if pid_file.is_file():
                pid = int(pid_file.read_text().strip())
                subprocess.run([shell, "-NoProfile", "-NonInteractive", "-Command",
                                f"Stop-Process -Id {pid} -ErrorAction SilentlyContinue"],
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0), timeout=10)
            return None
        mapping = json.loads(output.read_text(encoding="utf-8-sig"))
        if not isinstance(mapping, dict) or mapping.get("source_sha256") != hashlib.sha256(Path(path).read_bytes()).hexdigest():
            return None
        return mapping
    except (OSError, ValueError, subprocess.SubprocessError):
        return None


def extract_docx(path: str, key: str = "", pagination=None) -> Extracted:
    """Paragraphs in document order, tables as pipe tables, and every
    embedded picture unpacked with a marker where it occurred."""
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(path)
    except Exception as exc:
        raise UnreadableFile(*_why(exc)) from exc

    blip_tag, embed_attr = qn("a:blip"), qn("r:embed")
    collector = _ImageCollector(os.path.basename(path), key)
    parts: list = []
    para_no = 0
    hierarchy = []
    numbered = {}
    numbering = _WordNumbering(document)

    def add_images(element, locator):
        for blip in element.iter(blip_tag):
            rid = blip.get(embed_attr)
            related = document.part.related_parts.get(rid) if rid else None
            if related is not None:
                img = collector.add(related.blob,
                                    os.path.splitext(str(related.partname))[1],
                                    locator)
                parts.append(marker_for(img.image_id))

    blocks = list(_docx_blocks(document.element.body))
    pages = _word_page_blocks(path, blocks, pagination) if pagination else None
    current_page = None
    for index, child in enumerate(blocks):
        page_info = pages[index] if pages else None
        if page_info and page_info["start_page"] != current_page:
            current_page = page_info["start_page"]
            parts.append(f"--- Page {current_page} ---")
            # Carry current chapter on continuation pages, so a page+chapter
            # locator has an explicit marker in that page's extracted region.
            is_new_heading = child.tag.rsplit("}", 1)[-1] == "p" and _is_heading(Paragraph(child, document))
            if hierarchy and not is_new_heading:
                parts.append(f"--- Section: {hierarchy[-1][1]} ---")
                if len(hierarchy) > 1:
                    parts.append("--- Section Path: " + " > ".join(
                        item[1] for item in hierarchy) + " ---")
        tag = child.tag.split("}")[-1]
        span_marker = None
        if page_info and page_info["end_page"] > page_info["start_page"]:
            label = "Paragraph" if tag == "p" else "Table"
            span_marker = f"[[{label} spans rendered pages {page_info['start_page']}-{page_info['end_page']}]]"
            if tag != "p":
                parts.append(span_marker)
        if tag == "p":
            para_no += 1
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            level = _heading_level(paragraph)
            label = numbering.label(paragraph)
            heading = " ".join(text.split())
            if label and level is not None and not re.match(r"^\d", heading):
                heading = label + " " + heading
            manual = _MANUAL_SECTION.match(heading)
            number = manual.group(1) if manual else ""
            # Decimal subdivisions are supported only under an observed
            # chapter/parent. This captures manually numbered RFQ headings
            # without mistaking decimal values or unrelated list items for
            # sections. Top-level numbers alone need outline/style evidence.
            if level is None and "." in number and len(heading) <= 180:
                parent = number.rsplit(".", 1)[0]
                # A decimal quantity or an ordinary requirement sentence is
                # body text, even when its integer matches a chapter number.
                body_sentence = re.search(
                    r"^\d+(?:\.\d+)+\s+(?:lit(?:re|er)s?|mm|cm|km|kg|kw|nm|bar|mpa|rpm|%)(?:\b|\s)|"
                    r"\b(?:shall|must|is required|are required)\b", heading, re.IGNORECASE)
                active_parent = next((item for item in reversed(hierarchy)
                                      if item[2] == parent), None)
                if not body_sentence and parent in numbered and active_parent and (
                        number.split(".")[0] == hierarchy[0][2].split(".")[0]):
                    level = len(number.split("."))
            if text:
                if level is not None:
                    while hierarchy and hierarchy[-1][0] >= level:
                        hierarchy.pop()
                    if level == 0:
                        hierarchy.clear()
                    else:
                        hierarchy.append((level, heading, number))
                    if number:
                        numbered[number] = heading
                    parts.append(f"--- Section: {heading} ---")
                    if len(hierarchy) > 1:
                        parts.append("--- Section Path: " + " > ".join(
                            item[1] for item in hierarchy) + " ---")
                    if span_marker:
                        parts.append(span_marker)
                else:
                    if span_marker:
                        parts.append(span_marker)
                    parts.append(text)
            locator = ("section " + hierarchy[-1][1] if hierarchy
                       else f"paragraph {para_no}")
            add_images(child, locator)
        elif tag == "tbl":
            table = Table(child, document)
            md = rows_to_markdown([_docx_row_cells(r) for r in table.rows])
            if md:
                parts.append(md)
            add_images(child, "section " + hierarchy[-1][1] if hierarchy
                       else f"table after paragraph {para_no}")

    locations = ({"status": "available", "method": pagination["method"],
                  "pages": pagination["pages"], "blocks": pages}
                 if pages else {"status": "unavailable", "method": "chapter headings",
                                "reason": "Word rendered pagination unavailable"})
    return Extracted("\n\n".join(parts), collector.images, locations)


# ------------------------------------------------------------------- pptx
def extract_pptx(path: str, key: str = "") -> Extracted:
    """Per slide: title, body text, tables as pipe tables, speaker notes."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        prs = Presentation(path)
    except Exception as exc:
        raise UnreadableFile(*_why(exc)) from exc

    collector = _ImageCollector(os.path.basename(path), key)
    parts: list = []

    for n, slide in enumerate(prs.slides, start=1):
        chunk = [f"--- Slide {n} ---"]
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        if title:
            chunk.append(f"Title: {title}")
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = collector.add(shape.image.blob,
                                    "." + str(shape.image.ext).lstrip("."),
                                    f"slide {n}")
                chunk.append(marker_for(img.image_id))
                continue
            if shape.has_table:
                md = rows_to_markdown([_pptx_row_cells(r)
                                       for r in shape.table.rows])
                if md:
                    chunk.append(md)
                continue
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text and text != title:
                    chunk.append(text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                chunk.append(f"Speaker notes: {notes}")
        parts.append("\n".join(chunk))

    return Extracted("\n\n".join(parts), collector.images)


# ------------------------------------------------------------------- xlsx
def extract_xlsx(path: str, key: str = "") -> Extracted:
    """Per sheet: header + rows as a pipe table, empty rows skipped."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise UnreadableFile(*_why(exc)) from exc

    parts: list = []
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    rows.append(cells)
            md = rows_to_markdown(rows)
            if md:
                parts.append(f"--- Sheet: {ws.title} ---\n{md}")
    except Exception as exc:
        raise UnreadableFile(*_why(exc)) from exc
    finally:
        wb.close()

    return Extracted("\n\n".join(parts), [])


# -------------------------------------------------------------------- pdf
# Render density for anything handed to the visual pass. 300 dpi is the
# long-standing floor for reading text off a scan, and these renders exist
# to be read: at 150 dpi the annotation text on a dense process-flow diagram
# is illegible, so the figure gets flagged, comes back transcribed as
# unreadable, and the flagging achieves nothing for exactly the figures
# worth the most. Measured on a real package: a module PFD whose equipment
# tags and stack ratings were unreadable at 150 and plain at 300.
RENDER_DPI = 300


def _page_png(page) -> bytes:
    """Render a page to PNG so a scanned page can be read visually."""
    buf = io.BytesIO()
    page.to_image(resolution=RENDER_DPI).original.save(buf, format="PNG")
    return buf.getvalue()


# A figure embedded in a text-bearing PDF page has to be worth a visual read
# before we cut it out. Expressed in PDF points (72 per inch), so this is
# about 0.8 x 0.14 inches — the size of a single line of small text.
#
# Deliberately set at LABEL size, not chart size, and the cost of that is
# accepted. Two findings from a real package drove it down here. Its only
# telephone number and the engineering firm's name sat inside a 188 x 42
# point letterhead strip. Worse, its seismic sheet rendered the standard's
# version as a 72 x 17 point PICTURE reading "ASCE 7-22" over a text layer
# that still said "ASCE 7-16" — a version number is exactly the identifier
# the analysis prompt requires be reproduced exactly, and a chart-sized
# floor silently reported the wrong one.
#
# This flags several times more figures per package than a chart-sized floor
# would. That is the intended trade: the knowledge pack names letterheads,
# signature blocks and footers as the first places contact details hide, and
# a missed contact or a wrong standard version is a reported error, while an
# extra transcription is only time.
MIN_PDF_IMAGE_POINTS = 60 * 10


def _page_region_png(page, bbox) -> bytes:
    """Render one region of a page to PNG.

    Cropping and rendering rather than pulling the embedded stream out: PDF
    images arrive in several encodings (DCTDecode, FlateDecode, CCITT) and
    some carry no standalone container, so reconstructing bytes per encoding
    is far more fragile than asking the renderer for the pixels.
    """
    x0, top, x1, bottom = bbox
    x0 = max(x0, page.bbox[0])
    top = max(top, page.bbox[1])
    x1 = min(x1, page.bbox[2])
    bottom = min(bottom, page.bbox[3])
    buf = io.BytesIO()
    page.crop((x0, top, x1, bottom)).to_image(resolution=RENDER_DPI) \
        .original.save(buf, format="PNG")
    return buf.getvalue()


def _widen_to_graphic_band(page, bbox):
    """Grow a figure's box sideways over adjoining artwork that holds no text.

    Some content is drawn as VECTOR PATHS rather than set as text or placed
    as a bitmap. A letterhead is the common case: the logo is a bitmap, and
    the company name, address and phone beside it are outlined shapes. Such
    a block is invisible three times over — extract_text finds no words,
    page.images does not list it, and cropping the bitmap alone cuts it off
    the right-hand edge. Only rendering reveals it, which is exactly how the
    OCR-based system this skill replaces used to read it.

    So the crop grows horizontally across the figure's own rows to take in
    neighbouring vector artwork, and stops the moment a real text word falls
    inside — body text is already extracted and must not be duplicated.
    """
    x0, top, x1, bottom = bbox
    px0, ptop, px1, pbottom = page.bbox
    band_top, band_bottom = top - 4, bottom + 4

    def in_band(o):
        try:
            return (float(o["bottom"]) > band_top
                    and float(o["top"]) < band_bottom)
        except (KeyError, TypeError, ValueError):
            return False

    grown_x0, grown_x1 = x0, x1
    grown_top, grown_bottom = top, bottom
    for obj in list(page.rects or []) + list(page.curves or []):
        if in_band(obj):
            try:
                grown_x0 = min(grown_x0, float(obj["x0"]))
                grown_x1 = max(grown_x1, float(obj["x1"]))
                # Vector lettering usually overshoots the bitmap's rows
                # slightly; without this the first and last lines of an
                # address block come out sliced in half.
                grown_top = min(grown_top, float(obj["top"]))
                grown_bottom = max(grown_bottom, float(obj["bottom"]))
            except (KeyError, TypeError, ValueError):
                continue

    # Never swallow extracted text: that content is already in the corpus,
    # and a transcription of it would duplicate rather than add.
    try:
        words = page.extract_words() or []
    except Exception:
        words = []
    for w in words:
        try:
            if (float(w["bottom"]) > grown_top and float(w["top"]) < grown_bottom
                    and float(w["x1"]) > grown_x0
                    and float(w["x0"]) < grown_x1):
                return (x0, top, x1, bottom)
        except (KeyError, TypeError, ValueError):
            continue

    return (max(grown_x0, px0), max(grown_top, ptop),
            min(grown_x1, px1), min(grown_bottom, pbottom))


def _embedded_page_images(page):
    """(boxes worth reading, count of figures below the floor).

    A page that HAS a text layer still hides content inside its figures —
    letterheads carrying a company name and phone number, charts, scanned
    tables, labels rendered as pictures. Native text extraction cannot see
    any of it, whereas OCR-the-whole-page (what this skill replaced) always
    could. Unpacking them here closes that regression.

    The count of figures below the floor is returned rather than discarded:
    those are usually icons and rules, but not always, and a silent loss is
    worse than a noted one.
    """
    found, skipped = [], 0
    for im in page.images or []:
        try:
            x0, top = float(im["x0"]), float(im["top"])
            x1, bottom = float(im["x1"]), float(im["bottom"])
        except (KeyError, TypeError, ValueError):
            continue
        if (x1 - x0) * (bottom - top) >= MIN_PDF_IMAGE_POINTS:
            found.append(_widen_to_graphic_band(page, (x0, top, x1, bottom)))
        else:
            skipped += 1
    # Reading order: top to bottom, then left to right. Deterministic, and it
    # keeps each marker near the text it sits beside.
    found.sort(key=lambda b: (round(b[1], 1), round(b[0], 1)))
    return found, skipped


def _pdf_outline_sections(pdf):
    """Resolve actual PDF outline destinations to physical page numbers."""
    from pdfminer.pdftypes import PDFObjRef, resolve1
    page_numbers = {p.page_obj.pageid: i for i, p in enumerate(pdf.pages, 1)}
    sections = {}
    hierarchy = []
    try:
        outlines = pdf.doc.get_outlines()
        for level, title, dest, action, _ in outlines:
            try:
                if dest is None and action is not None:
                    dest = resolve1(action).get("D")
                if isinstance(dest, (str, bytes)):
                    dest = pdf.doc.get_dest(dest)
                dest = resolve1(dest)
                if isinstance(dest, dict):
                    dest = resolve1(dest.get("D"))
                reference = dest[0]
                page = page_numbers.get(reference.objid) if isinstance(reference, PDFObjRef) else None
                heading = " ".join(str(title).split())
                if page is None or not heading:
                    continue
                while hierarchy and hierarchy[-1][0] >= level:
                    hierarchy.pop()
                hierarchy.append((level, heading))
                sections.setdefault(page, []).append(
                    (heading, " > ".join(h for _, h in hierarchy)))
            except (AttributeError, IndexError, KeyError, TypeError, ValueError):
                continue
    except Exception:
        # A missing/broken outline does not make the PDF's text unreadable.
        return {}
    return sections


def _pdf_heading_text(page, text, hierarchy, tables=None):
    """Conservative numbered headings backed by visible PDF typography.

    A top chapter must be larger and bold. Decimal descendants need an
    observed chapter, bold heading typography, and a matching number prefix.
    Exclude table regions and TOC dot leaders; never treat body list items,
    decimal measurements, or an ordinary numbered sentence as a heading.
    """
    try:
        lines = page.extract_text_lines()
        chars = page.chars
        sizes = [round(float(c["size"]), 1) for c in chars if c.get("text", "").strip()]
        if not sizes:
            return text
        from collections import Counter
        body_size = Counter(sizes).most_common(1)[0][0]
        boxes = [t.bbox for t in (page.find_tables() if tables is None else tables)]
    except (AttributeError, KeyError, TypeError, ValueError):
        return text
    numbered = {h[2]: h[1] for h in hierarchy if h[2]}
    replacements = {}
    consumed = set()
    for index, line in enumerate(lines):
        if index in consumed:
            continue
        heading = " ".join(line.get("text", "").split())
        first_line = heading
        match = _MANUAL_SECTION.match(heading)
        if not match or len(heading) > 180 or re.search(r"\.{3,}|\s\d+\s*$", heading):
            continue
        glyphs = [c for c in line.get("chars", []) if c.get("text", "").strip()]
        if not glyphs or any(b[0] <= line["x0"] < b[2] and b[1] <= line["top"] < b[3]
                             for b in boxes):
            continue
        bold = sum(bool(re.search(r"bold|black|demi", c.get("fontname", ""), re.I))
                   for c in glyphs) / len(glyphs)
        number = match.group(1)
        level = len(number.split("."))
        if bold < 0.8:
            continue
        if level == 1:
            if max(float(c["size"]) for c in glyphs) <= body_size + 1:
                continue
        elif (not any(h[2] == number.rsplit(".", 1)[0] for h in hierarchy) or not hierarchy
              or number.split(".")[0] != hierarchy[0][2].split(".")[0]):
            continue
        # Retain wrapped heading words only when their typography and nearby
        # placement corroborate a continuation, never swallow normal body.
        previous = line
        for next_index in range(index + 1, min(index + 3, len(lines))):
            continuation = lines[next_index]
            words = " ".join(continuation.get("text", "").split())
            more = [c for c in continuation.get("chars", []) if c.get("text", "").strip()]
            if (not words or _MANUAL_SECTION.match(words) or not more
                    or continuation["top"] - previous["bottom"] > 24
                    or abs(continuation["x0"] - line["x0"]) > 30
                    or sum(bool(re.search(r"bold|black|demi", c.get("fontname", ""), re.I))
                           for c in more) / len(more) < 0.8
                    or max(float(c["size"]) for c in more) < body_size
                    or len(heading + words) > 180):
                break
            heading += " " + words
            replacements[words] = ""
            consumed.add(next_index)
            previous = continuation
        while hierarchy and hierarchy[-1][0] >= level:
            hierarchy.pop()
        hierarchy.append((level, heading, number))
        numbered[number] = heading
        marker = f"--- Section: {heading} ---"
        if len(hierarchy) > 1:
            marker += "\n--- Section Path: " + " > ".join(h[1] for h in hierarchy) + " ---"
        replacements[first_line] = marker
    return "\n".join(replacements.get(" ".join(line.split()), line)
                     for line in text.splitlines())


def _pdf_table_text(page, text, tables):
    """Remove only glyphs accounted for, in order, by an extracted cell.

    A missing word, changed qualifier, overlapping cell or unassigned glyph
    keeps native text. Geometry binds a representation to its own table;
    repetition elsewhere on the page is never a reason to remove anything.
    """
    from pdfplumber.utils import extract_text

    def compact(value):
        return re.sub(r"\s+", " ", value or "").strip()

    def inside(glyph, box):
        return (box[0] - .01 <= glyph["x0"] and glyph["x1"] <= box[2] + .01
                and box[1] - .01 <= glyph["top"] and glyph["bottom"] <= box[3] + .01)

    rendered, verified, removed = [], [], set()
    for table in tables:
        try:
            rows = table.extract()
            rendered.append((table, rows, rows_to_markdown(rows)))
        except Exception:
            # A malformed native table must not discard readable page text.
            continue
    original = [md for _, _, md in rendered if md]
    try:
        native_lines = page.extract_text_lines(keep_blank_chars=True)
        native_words = page.extract_words()
        for table, rows, md in rendered:
            if not md or len(rows) != len(table.rows):
                continue
            # A token straddling a boundary cannot be split between native
            # prose and a table cell without changing its meaning/heading.
            if any(c.get("text", "").strip() and
                   min(c["x1"], table.bbox[2]) > max(c["x0"], table.bbox[0]) and
                   min(c["bottom"], table.bbox[3]) > max(c["top"], table.bbox[1]) and
                   not inside(c, table.bbox) for c in page.chars):
                continue
            if any(other is not table and min(table.bbox[2], other.bbox[2]) > max(table.bbox[0], other.bbox[0])
                   and min(table.bbox[3], other.bbox[3]) > max(table.bbox[1], other.bbox[1])
                   for other in tables) or sum(other is table for other in tables) > 1:
                continue
            if any(min(word["x1"], table.bbox[2]) > max(word["x0"], table.bbox[0]) and
                   min(word["bottom"], table.bbox[3]) > max(word["top"], table.bbox[1]) and
                   sum(inside(word, box) for box in table.cells if box is not None) != 1
                   for word in native_words):
                continue
            # A heading/prose line beside the first row makes the chapter
            # ordering ambiguous. Preserve native order instead of moving
            # the entire table ahead of that source context.
            first_top = min((c["top"] for c in page.chars if inside(c, table.bbox)
                             and c.get("text", "").strip()), default=table.bbox[1])
            if any(abs(line["top"] - first_top) <= 12 and
                   any(c.get("text", "").strip() and not inside(c, table.bbox)
                       for c in line.get("chars", [])) for line in native_lines):
                continue
            cells, complete = [], True
            for geometry, values in zip(table.rows, rows):
                if len(geometry.cells) != len(values):
                    complete = False
                    break
                for box, value in zip(geometry.cells, values):
                    if box is None:
                        if compact(value):
                            complete = False
                        continue
                    glyphs = [c for c in page.chars if inside(c, box)]
                    if compact(extract_text(glyphs) or "") != compact(value):
                        complete = False
                        break
                    cells.append((box, glyphs))
                if not complete:
                    break
            glyphs = [c for c in page.chars if inside(c, table.bbox)]
            substantive = [c for c in glyphs if c.get("text", "").strip()]
            # Includes merged cells only when their actual bounding boxes
            # account for every glyph exactly once. Unknown merges fall back.
            if not substantive or any(sum(inside(c, box) for box, _ in cells) != 1
                                      for c in substantive):
                complete = False
            if complete and not any(id(c) in removed for c in substantive):
                verified.append((table, md))
                removed.update(id(c) for c in glyphs)
        if not verified:
            return text, original, 0, original
        lines = native_lines
        line_chars = {id(c) for line in lines for c in line.get("chars", [])}
        if any(id(c) not in line_chars for c in page.chars if c.get("text", "").strip()):
            return text, original, 0, original
        pending = sorted(verified, key=lambda pair: (pair[0].bbox[1], pair[0].bbox[0]))
        pieces = []
        for line in lines:
            while pending and pending[0][0].bbox[1] <= line["top"]:
                pieces.append(pending.pop(0)[1])
            keep = [c for c in line.get("chars", []) if id(c) not in removed]
            # Keep literal whitespace and the original geometry's inferred
            # word gaps. Untouched lines retain their exact extracted text.
            value = line["text"] if len(keep) == len(line.get("chars", [])) else (
                (extract_text(keep, keep_blank_chars=True) or "").strip())
            if value:
                pieces.append(value)
        pieces.extend(md for _, md in pending)
        fallback = [md for table, _, md in rendered if md and
                    not any(table is good for good, _ in verified)]
        return "\n".join(pieces), fallback, len(verified), original
    except Exception:
        # Preserve all native evidence when a library lacks reliable geometry.
        return text, original, 0, original


def extract_pdf(path: str, key: str = "") -> Extracted:
    """Text layer plus extract_tables() per page. A page with no text layer
    is a scan: it is rendered to PNG and marked for a visual read, which is
    what the old OCR-the-whole-PDF path used to cover."""
    import pdfplumber

    collector = _ImageCollector(os.path.basename(path), key)
    parts: list = []
    raw_parts: list = []
    verified_tables = 0
    try:
        with pdfplumber.open(path) as pdf:
            outlines = _pdf_outline_sections(pdf)
            hierarchy = []
            current_outline = []
            for n, page in enumerate(pdf.pages, start=1):
                chunk = [f"--- Page {n} ---"]
                text = (page.extract_text() or "").strip()
                try:
                    tables = page.find_tables() or []
                except Exception:
                    tables = []
                optimized, fallback, count, original_tables = _pdf_table_text(page, text, tables)
                verified_tables += count
                raw_chunk = list(chunk)
                if outlines:
                    if n in outlines:
                        current_outline = outlines[n]
                    for heading, section_path in current_outline:
                        chunk.append(f"--- Section: {heading} ---")
                        if section_path != heading:
                            chunk.append(f"--- Section Path: {section_path} ---")
                    raw_chunk = list(chunk)
                    # All destinations on their first page are valid; only
                    # the last active chapter continues onto following pages.
                    if current_outline:
                        current_outline = current_outline[-1:]
                elif text:
                    initial = list(hierarchy)
                    text = _pdf_heading_text(page, text, hierarchy, tables)
                    optimized = _pdf_heading_text(page, optimized, initial, tables)
                if optimized:
                    chunk.append(optimized)
                chunk.extend(md for md in fallback if md)
                if text:
                    raw_chunk.append(text)
                raw_chunk.extend(md for md in original_tables if md)
                marker_start = len(chunk)
                if not text:
                    try:
                        img = collector.add(_page_png(page), ".png",
                                            f"page {n}")
                        chunk.append(marker_for(img.image_id))
                    except Exception:
                        chunk.append(f"[[PAGE {n} could not be read]]")
                else:
                    # The page has text, but its figures do not. Each one is
                    # cut out and marked so the visual pass can read it.
                    boxes, skipped = _embedded_page_images(page)
                    for bbox in boxes:
                        try:
                            img = collector.add(_page_region_png(page, bbox),
                                                ".png", f"page {n}",
                                                presized=True)
                        except Exception:
                            continue
                        chunk.append(marker_for(img.image_id))
                    if skipped:
                        # Say so rather than lose it silently. Small figures
                        # are usually icons, but a label rendered as an image
                        # can carry a standard's version number — one real
                        # package showed "ASCE 7-22" as a picture over a text
                        # layer still reading "ASCE 7-16".
                        chunk.append(
                            f"[[{skipped} small figure(s) on this page were "
                            f"not read; any text inside them is missing]]")
                parts.append("\n".join(chunk))
                raw_chunk.extend(chunk[marker_start:])
                raw_parts.append("\n".join(raw_chunk))
    except Exception as exc:
        raise UnreadableFile(*_why(exc)) from exc

    optimized = "\n\n".join(parts)
    raw = "\n\n".join(raw_parts)
    return Extracted(optimized, collector.images, raw_text=raw,
                     extraction_metrics={"verified_tables": verified_tables,
                                         "raw_chars": len(raw),
                                         "model_chars": len(optimized),
                                         "removed_chars": len(raw) - len(optimized)})


# ------------------------------------------------------------- text / csv
def extract_text_file(path: str, key: str = "") -> Extracted:
    """.txt and .csv, decoded and normalized to \n line endings.

    Encodings are tried in order of likelihood. Line endings are normalized
    because a corpus that mixes CRLF and LF makes every downstream character
    count and diff platform-dependent for no benefit.
    """
    try:
        with open(path, "rb") as fh:
            raw = fh.read()
    except Exception as exc:
        raise UnreadableFile(*_why(exc)) from exc
    for encoding in _ENCODINGS:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    return Extracted(text.replace("\r\n", "\n").replace("\r", "\n"), [])


# --------------------------------------------------------------- dispatch
_DISPATCH = {
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
    ".pdf": extract_pdf,
    ".csv": extract_text_file,
    ".txt": extract_text_file,
}


def extract_file(path: str, key: str = "") -> Extracted:
    """Extract one file. Raises UnreadableFile with a plain-language reason.

    `key` is the corpus-unique identity every image found here is stamped
    with (see image_key). It defaults to the basename, which is correct
    whenever the corpus holds no two files of that name.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnreadableFile("unsupported file type")
    if ext in IMAGE_EXTENSIONS:
        try:
            with open(path, "rb") as fh:
                data = fh.read()
        except Exception as exc:
            raise UnreadableFile(*_why(exc)) from exc
        name = os.path.basename(path)
        return Extracted("", [RawImage(name, 1, data, ext, "standalone",
                                       key or name)])
    return _DISPATCH[ext](path, key)


# Pattern-based cleanup of native document extraction artifacts.
_PAGE_NUM = re.compile(r"^\s*(page\s+)?\d{1,4}(\s*(of|/)\s*\d{1,4})?\s*$",
                       re.IGNORECASE)

# classification banners: a short line consisting (almost) only of a
# company name and/or a classification word
_CLASS_WORDS = (r"(restricted|confidential|internal( use)?( only)?|proprietary"
                r"|uncontrolled( copy| when printed)?|company confidential"
                r"|all rights reserved)")
_BANNER = re.compile(
    r"^\s*(\S+\s+)?" + _CLASS_WORDS + r"(\s+\S+)?\s*$", re.IGNORECASE)
_BANNER_MAX_LEN = 45


def clean_extracted_text(text: str) -> str:
    """Clean one file's extracted_text. Safe to call on anything."""
    if not text:
        return text
    # normalize whitespace variants the extractor emits
    text = text.replace(" ", " ").replace("\t", " ")
    out = []
    for raw in text.split("\n"):
        s = raw.strip()
        if _PAGE_NUM.match(s):
            continue
        if len(s) <= _BANNER_MAX_LEN and _BANNER.match(s):
            continue
        out.append(raw.rstrip())
    cleaned = "\n".join(out)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)   # collapse blank-line runs
    return cleaned.strip()


def clean_corpus_files(files: list[dict]) -> tuple[list[dict], dict]:
    """Clean every file's extracted_text; returns (files, stats) — stats go
    into the run payload so the effect is auditable."""
    stats = {}
    cleaned_files = []
    for f in files:
        original = f.get("extracted_text", "") or ""
        cleaned = clean_extracted_text(original)
        g = dict(f)
        g["extracted_text"] = cleaned
        cleaned_files.append(g)
        stats[f.get("filename", "?")] = {
            "chars_before": len(original),
            "chars_after": len(cleaned),
            "removed_pct": round(100 * (1 - len(cleaned) / len(original)), 1)
            if original else 0.0,
        }
    return cleaned_files, stats
